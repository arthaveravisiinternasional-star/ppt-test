"""Build the SCANVAS GESAwards 2026 pitch deck from the Freesia template.

    python build_deck.py                 # -> output/SCANVAS_GESAwards2026.pptx
    python build_deck.py --render        # + PNG previews in output/preview (needs PowerPoint)
    python build_deck.py --render 3 7    # preview only slides 3 and 7
    python build_deck.py --pdf           # embed Fredoka + Varela Round into the PPTX, then export PDF
"""
import argparse
import subprocess
import sys

from pptx import Presentation

from deck import recolor
from deck import theme as T
from deck.slides import SLIDES

OUT = T.OUTPUT / 'SCANVAS_GESAwards2026.pptx'


def build(out=OUT):
    prs = Presentation(str(T.TEMPLATE))
    recolor.apply(prs)                  # Freesia colors -> Scanvas DESIGN.md palette
    layouts = {L.name: L for m in prs.slide_masters for L in m.slide_layouts}
    ids = prs.slides._sldIdLst
    for sid in list(ids):              # drop the 50 sample slides, keep masters/layouts
        prs.part.drop_rel(sid.rId)
        ids.remove(sid)
    for n, fn in enumerate(SLIDES, 1):
        s = prs.slides.add_slide(layouts[T.LAYOUT[n]])
        for ph in list(s.placeholders):  # keep only the title placeholder
            if ph.placeholder_format.idx != 0:
                ph._element.getparent().remove(ph._element)
        fn(s)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    print('saved', out)
    return out


def powerpoint(out, mode, slides=()):
    script = T.ROOT / 'tools' / 'render.ps1'
    cmd = ['powershell', '-NoProfile', '-ExecutionPolicy', 'Bypass', '-File', str(script),
           '-File', str(out), '-Mode', mode, '-Dir', str(T.OUTPUT / 'preview')]
    if slides:
        cmd += ['-Slides', ','.join(map(str, slides))]
    subprocess.run(cmd, check=True)


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--render', nargs='*', type=int, help='export PNG previews (optionally only these slide numbers)')
    ap.add_argument('--pdf', action='store_true')
    ap.add_argument('--out', help='output .pptx path (default output/SCANVAS_GESAwards2026.pptx)')
    a = ap.parse_args()
    out = build(T.ROOT / a.out if a.out else OUT)
    if a.render is not None:
        powerpoint(out, 'png', a.render)
    if a.pdf:                          # final deliverable: embed fonts, then export PDF
        powerpoint(out, 'embed')
        powerpoint(out, 'pdf')
    sys.exit(0)
