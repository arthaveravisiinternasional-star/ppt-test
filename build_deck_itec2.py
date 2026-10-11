"""Build the ITEC SEA 2026 Top 30 (v2) edition: 10 main slides + 3 appendix, Digital Culture Track.
Backbone = ITEC v1; evidence from DSC, growth story from Antler, future loop from GESA. Other editions untouched.

    python build_deck_itec2.py                  # -> output/itec2/SCANVAS_ITEC_SEA2026_v2.pptx
    python build_deck_itec2.py --render 6 7     # previews in output/itec2/preview
    python build_deck_itec2.py --render --pdf   # final: embed fonts + PDF
"""
import argparse
import subprocess
import sys

from pptx import Presentation

import deck
from deck import content_itec2

sys.modules['deck.content'] = content_itec2   # swap copy before slide modules import it
deck.content = content_itec2

from deck import recolor, slides as S, slides_dsc as D, slides_itec2 as I  # noqa: E402
from deck import theme as T  # noqa: E402

OUT_DIR = T.OUTPUT / 'itec2'
OUT = OUT_DIR / 'SCANVAS_ITEC_SEA2026_v2.pptx'

DECK = [(S.s1_cover, 'TITLE'),                      # 1 Cover
        (I.s_problem, 'TITLE_ONLY_3_2'),            # 2 Problem
        (S.s4_solution, 'TITLE_ONLY'),              # 3 How it works (no bottom strip: slide rules)
        (I.s_live, 'TITLE_ONLY_3_1_1_1'),           # 4 Already live
        (I.s_tech, 'TITLE_ONLY_2_1'),               # 5 Technology / defensibility
        (I.s_culture, 'TITLE_ONLY'),                # 6 Why Digital Culture
        (I.s_market, 'TITLE_AND_BODY_1_1'),         # 7 Market: wedge -> scale -> expand
        (I.s_business, 'TITLE_AND_BODY'),           # 8 Business model + GTM
        (D.s_traction, 'TITLE_ONLY_3_1_2'),         # 9 Traction
        (I.s_next, 'BLANK_1_1'),                    # 10 Roadmap + team
        (D.s_competition, 'TITLE_ONLY'),            # A Competitive landscape
        (D.s_sources, 'TITLE_ONLY_3_1_2'),          # B Market evidence
        (I.s_pipeline, 'TITLE_ONLY_3_2')]           # C Processing pipeline


def build(out=OUT):
    prs = Presentation(str(T.TEMPLATE))
    recolor.apply(prs)
    layouts = {L.name: L for m in prs.slide_masters for L in m.slide_layouts}
    ids = prs.slides._sldIdLst
    for sid in list(ids):
        prs.part.drop_rel(sid.rId)
        ids.remove(sid)
    for fn, layout in DECK:
        s = prs.slides.add_slide(layouts[layout])
        for ph in list(s.placeholders):
            if ph.placeholder_format.idx != 0:
                ph._element.getparent().remove(ph._element)
        fn(s)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    print('saved', out)
    return out


def powerpoint(out, mode, slides=()):
    cmd = ['powershell', '-NoProfile', '-ExecutionPolicy', 'Bypass', '-File', str(T.ROOT / 'tools' / 'render.ps1'),
           '-File', str(out), '-Mode', mode, '-Dir', str(OUT_DIR / 'preview')]
    if slides:
        cmd += ['-Slides', ','.join(map(str, slides))]
    subprocess.run(cmd, check=True)


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--render', nargs='*', type=int)
    ap.add_argument('--pdf', action='store_true')
    ap.add_argument('--out', help='output .pptx path relative to project root')
    a = ap.parse_args()
    out = build(T.ROOT / a.out if a.out else OUT)
    if a.render is not None:
        powerpoint(out, 'png', a.render)
    if a.pdf:
        powerpoint(out, 'embed')
        powerpoint(out, 'pdf')
