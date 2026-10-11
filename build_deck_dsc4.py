"""Build DSC Season 17 v4 (Oct 2026): same template, look and slide order as DSC v3 (15 main + 4 appendix), with the
ITEC v3 validation story and fuller commercial/financial evidence. DSC v3 files and output are not touched.

    python build_deck_dsc4.py                  # -> output/dsc4/SCANVAS_DSC_S17_v4.pptx
    python build_deck_dsc4.py --render 5 10    # previews in output/dsc4/preview
    python build_deck_dsc4.py --render --pdf   # final: embed fonts + PDF
"""
import argparse
import subprocess
import sys

from pptx import Presentation

import deck
from deck import content_dsc4

sys.modules['deck.content'] = content_dsc4  # swap copy before slide modules import it
deck.content = content_dsc4

from deck import recolor, slides as S, slides_dsc as D, slides_dsc4 as V, slides_itec3 as I  # noqa: E402
from deck import theme as T  # noqa: E402

OUT_DIR = T.OUTPUT / 'dsc4'
OUT = OUT_DIR / 'SCANVAS_DSC_S17_v4.pptx'

# (slide function, Freesia layout) in DSC guideline order
DECK = [(S.s1_cover, 'TITLE'),                      # 1 Cover                       keep
        (D.s_overview, 'TITLE_ONLY_3_1_1_1'),       # 2 Business Overview           minor
        (V.s_problem, 'TITLE_ONLY_3_2'),            # 3 Problem & Opportunity       minor
        (D.s_solution, 'TITLE_ONLY'),               # 4 Solution                    keep
        (I.s_market, 'TITLE_AND_BODY_1_1'),         # 5 Market Size & Target Market major
        (I.s_competition, 'TITLE_ONLY'),            # 6 Competitive Landscape       major
        (V.s_business, 'TITLE_AND_BODY'),           # 7 Business Model              major
        (I.s_traction, 'TITLE_ONLY_3_1_2'),         # 8 Traction & Validation       major
        (V.s_gtm, 'TITLE_ONLY'),                    # 9 Marketing & GTM             minor
        (D.s_impact, 'TITLE_ONLY_2_1'),             # 10 ESG / SDGs / Impact        keep
        (V.s_financial, 'TITLE_AND_BODY'),          # 11 Financial Overview         major
        (V.s_funding, 'TITLE_ONLY_3_2'),            # 12 Funding Needs              major
        (S.s12_team, 'SECTION_HEADER_1_1'),         # 13 Team                       minor
        (V.s_roadmap, 'BLANK_1_1'),                 # 14 Milestones & Roadmap       major
        (V.s_closing, 'CAPTION_ONLY'),              # 15 Closing                    minor
        (S.s7_different, 'TITLE_ONLY_2_1'),         # A Technology edge
        (D.s_sources, 'TITLE_ONLY_3_1_2'),          # B Market data
        (V.s_pnl, 'TITLE_ONLY_3_2'),                # C 12-month P&L and cash
        (D.s_competitors, 'TITLE_ONLY_3_1_2')]      # D Competitor evidence


def build(out=OUT):
    prs = Presentation(str(T.TEMPLATE))
    recolor.apply(prs)
    layouts = {L.name: L for m in prs.slide_masters for L in m.slide_layouts}
    ids = prs.slides._sldIdLst
    for sid in list(ids):
        prs.part.drop_rel(sid.rId)
        ids.remove(sid)
    for fn, layout in DECK:
        s = prs.slides.add_slide(layouts[layout])
        for ph in list(s.placeholders):
            if ph.placeholder_format.idx != 0:
                ph._element.getparent().remove(ph._element)
        fn(s)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    print('saved', out)
    return out


def powerpoint(out, mode, slides=()):
    cmd = ['powershell', '-NoProfile', '-ExecutionPolicy', 'Bypass', '-File', str(T.ROOT / 'tools' / 'render.ps1'),
           '-File', str(out), '-Mode', mode, '-Dir', str(OUT_DIR / 'preview')]
    if slides:
        cmd += ['-Slides', ','.join(map(str, slides))]
    subprocess.run(cmd, check=True)


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--render', nargs='*', type=int)
    ap.add_argument('--pdf', action='store_true')
    ap.add_argument('--out', help='output .pptx path relative to project root')
    a = ap.parse_args()
    out = build(T.ROOT / a.out if a.out else OUT)
    if a.render is not None:
        powerpoint(out, 'png', a.render)
    if a.pdf:
        powerpoint(out, 'embed')
        powerpoint(out, 'pdf')
