"""Slides for the ITEC SEA 2026 official-template edition (v4). Coordinates stay on the 10 x 5.625 in grid; helpers
scale them 2x onto the committee's 20 x 11.25 in canvas. Keep content between y 0.3 and 5.3 (template frame bars).
Imported after build_deck_itec4.py swaps deck.content for deck.content_itec4."""
from PIL import Image
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

from . import content as K
from . import slides_dsc as D
from . import slides_itec3 as I3
from . import theme as T
from .helpers import (L, C, MID, R, text, title, rect, card, oval, pic, icon, icon_circle, chip, arrow_chips, notes)
from .slides_itec3 import _flow

ITEC_BLUE, ITEC_GREEN = '5271FF', '7ED957'   # template frame colors


def section(s, label):
    """Committee section name, top-right, in the template's blue."""
    text(s, 5.6, 0.33, 3.8, 0.2, label, size=7.5, font=T.LABEL_FONT, color=ITEC_BLUE, align=R, anchor=MID)


def _alpha(sh, val):
    clr = sh._element.spPr.find(qn('a:solidFill'))[0]
    clr.append(clr.makeelement(qn('a:alpha'), {'val': str(val)}))
    return sh


def s_cover(s):
    c = K.COVER
    pic(s, T.LOGO, 0.72, 0.6, h=0.5, alt='Scanvas logo')
    text(s, 0.72, 1.24, 4.0, 0.24, c['label'] + '  ·  ' + K.TRACK.upper(), size=8.5, font=T.LABEL_FONT,
         color=ITEC_BLUE, anchor=MID)
    title(s, c['title'], 0.7, 1.5, 4.2, 0.75, size=46)
    text(s, 0.72, 2.27, 3.9, 0.7, c['subtitle'], size=18, color=T.PURPLE, font=T.TITLE_FONT, spacing=0.95)
    text(s, 0.72, 3.0, 3.75, 0.6, c['summary'], size=11.5, color=T.BODY, spacing=1.12)
    y = 3.78
    for name, role in c['presenters']:
        text(s, 0.72, y, 3.9, 0.24, [[(name, {'font': T.LABEL_FONT, 'size': 10.5}),
                                      ('  ·  ' + role, {'size': 9, 'color': T.BODY})]], anchor=MID)
        y += 0.27
    text(s, 0.72, y + 0.02, 3.9, 0.22, c['company'], size=9, font=T.LABEL_FONT, color=T.PURPLE, anchor=MID)
    rect(s, 0.72, 4.82, 4.05, 0.3, fill=T.LAV2, radius=0.15)
    text(s, 0.72, 4.82, 4.05, 0.3, c['proof'], size=7.5, font=T.LABEL_FONT, color=T.DEEP, align=C, anchor=MID)
    # hero: colored paper -> alive on the tablet (ITEC logo from the template sits top-right)
    pic(s, 'paper_whale_tilt.png', 4.75, 0.5, w=2.2, alt='Child-colored Scanvas whale sheet')
    pic(s, 'tab_result.png', 5.0, 2.45, w=4.4, alt='The same whale scanned and alive in Scanvas')
    notes(s, c['notes'])


def s_product(s):
    """slides.s5_product re-spaced to clear the template's top bar, plus the TRY IT LIVE QR."""
    c = K.SLIDE_5_PRODUCT
    chip(s, c['tag'], 0.6, 0.5, bg=T.MINT, fg=T.GREEN, size=8.5, h=0.3, ic=('check', 'green'))
    title(s, c['title'], 0.6, 0.92, 2.85, 0.85, size=22)
    for i, (ic, head, detail) in enumerate(c['capabilities']):
        y = 2.0 + i * 0.52
        icon_circle(s, ic, 0.6, y + 0.02, 0.4)
        text(s, 1.12, y - 0.01, 3.1, 0.24, head, size=10.5, font=T.LABEL_FONT)
        text(s, 1.12, y + 0.22, 3.1, 0.24, detail, size=9.5, color=T.BODY)
    rect(s, 3.45, 0.62, 0.86, 0.86, fill=T.WHITE, line=T.BORDER, radius=0.1, lw=0.75)
    qr, _, _ = pic(s, T.ASSETS / K.SLIDE_4_SOLUTION['qr'], 3.5, 0.67, w=0.76, h=0.76,
                   alt='QR code to the Scanvas digital trial')
    qr.click_action.hyperlink.address = K.SLIDE_4_SOLUTION['qr_url']
    chip(s, c['qr_label'], 3.88, 1.54, bg=T.PURPLE, fg=T.WHITE, size=7, h=0.22, align='c')
    for i, (img, lbl) in enumerate(c['phones']):
        x = 4.5 + i * 1.62
        pic(s, img, x, 0.5, h=2.85, alt=lbl)
        text(s, x, 3.38, 1.45, 0.2, lbl, size=8.5, font=T.LABEL_FONT, color=T.BODY, align=C)
    for i, (img, lbl) in enumerate(c['screens']):
        x = 4.5 + i * 2.48
        pic(s, img, x, 3.68, w=2.4, alt=lbl)
        chip(s, lbl, x + 0.2, 3.82, bg=T.WHITE, fg=T.INK, size=7.5, h=0.24)
    notes(s, c['notes'])


def s_demo(s):
    """The digital trial as a video, with its five steps alongside."""
    from .helpers import Inches, chip_width, num_circle
    c = K.SLIDE_DEMO
    chip(s, c['tag'], 0.6, 0.5, bg=T.MINT, fg=T.GREEN, size=8.5, h=0.3, ic=('check', 'green'))
    title(s, c['title'], 0.6, 0.92, 2.7, 0.85, size=22)
    for i, (head, detail) in enumerate(c['steps']):
        y = 1.98 + i * 0.5
        num_circle(s, i + 1, 0.6, y + 0.02, 0.32, size=10)
        text(s, 1.02, y - 0.02, 2.2, 0.22, head, size=10, font=T.LABEL_FONT)
        text(s, 1.02, y + 0.19, 2.2, 0.22, detail, size=8.5, color=T.BODY)
    x, y, w = 3.3, 0.62, 6.15
    h = w * 9 / 16
    rect(s, x - 0.05, y - 0.05, w + 0.1, h + 0.1, fill=T.WHITE, line=T.BORDER, radius=0.06, lw=0.75)
    movie = s.shapes.add_movie(str(T.ASSETS / c['video']), Inches(x), Inches(y), Inches(w), Inches(h),
                               poster_frame_image=str(T.ASSETS / c['poster']), mime_type='video/mp4')
    movie.name = 'Scanvas digital trial demo'
    text(s, x, y + h + 0.12, w, 0.22, c['caption'], size=8.5, color=T.BODY, align=C, anchor=MID)
    cx = 0.6
    for label in c['chips']:
        chip(s, label, cx, 4.62, bg=T.LAV2, fg=T.DEEP, size=8, h=0.26)
        cx += chip_width(label, 8, h=0.26) + 0.08
    notes(s, c['notes'])


def s_about(s):
    c = K.SLIDE_ABOUT
    title(s, c['title'], 0.6, 0.45, 6.0, 0.8, size=22)
    text(s, 0.6, 1.35, 5.45, 0.8, c['summary'], size=11.5, color=T.BODY, spacing=1.15)
    rect(s, 0.6, 2.2, 5.45, 0.5, fill=T.LAV2, radius=0.14)
    text(s, 0.8, 2.2, 5.1, 0.5, c['mission'], size=10, font=T.LABEL_FONT, color=T.DEEP, anchor=MID, spacing=1.0)
    w, h = 2.67, 1.0
    for i, (ic, head, detail) in enumerate(c['facts']):
        x, y = 0.6 + (i % 2) * (w + 0.11), 2.88 + (i // 2) * (h + 0.12)
        card(s, x, y, w, h, radius=0.14)
        icon_circle(s, ic, x + 0.15, y + 0.15, 0.38)
        text(s, x + 0.62, y + 0.15, w - 0.75, 0.38, head, size=11, font=T.TITLE_FONT, anchor=MID)
        text(s, x + 0.15, y + 0.6, w - 0.3, 0.38, detail, size=8.5, color=T.BODY, spacing=1.05)
    oval(s, 6.55, 1.45, 2.75, fill=T.LAV2)
    pic(s, c['image'], 6.85, 0.62, h=4.45, alt='A child’s scanned whale swimming in the Scanvas Ocean world')
    notes(s, c['notes'])


def s_target(s):
    c = K.SLIDE_TARGET
    title(s, c['title'], 0.6, 0.45, 8.0, 0.5, size=24)
    w, top, h = 4.3, 1.2, 3.5
    for i, (tag, head, img, rows, flow, flow_label) in enumerate(c['markets']):
        x = 0.6 + i * (w + 0.2)
        card(s, x, top, w, h, radius=0.18)
        chip(s, tag, x + 0.2, top + 0.16, bg=T.PURPLE if i == 0 else T.MINT, fg=T.WHITE if i == 0 else T.GREEN,
             size=7.5, h=0.26)
        text(s, x + 0.2, top + 0.56, 2.3, 0.4, head, size=18, font=T.TITLE_FONT)
        pic(s, img, x + w - 1.72, top + 0.52, w=1.52, h=0.76, alt=f'{head} use case')
        for j, (lbl, v) in enumerate(rows):
            y = top + 1.36 + j * 0.54
            text(s, x + 0.2, y, w - 0.4, 0.17, lbl, size=7.5, font=T.LABEL_FONT, color=T.PURPLE)
            text(s, x + 0.2, y + 0.17, w - 0.4, 0.34, v, size=9, color=T.INK, spacing=1.0)
        text(s, x + 0.2, top + 2.97, w - 0.4, 0.17, flow_label, size=7.5, font=T.LABEL_FONT, color=T.PURPLE)
        _flow(s, ' → '.join(f.upper() for f in flow), x + 0.15, top + 3.15, w - 0.3, h=0.28, size=8)
    rect(s, 0.6, 4.84, 8.8, 0.36, fill=T.LAV2, radius=0.18)
    icon(s, 'palette', 'purple', 3.2, 4.92, 0.2)
    text(s, 3.45, 4.84, 4.5, 0.36, c['user'], size=10, font=T.TITLE_FONT, color=T.DEEP, anchor=MID)
    notes(s, c['notes'])


def s_market_size(s):
    c = K.SLIDE_MARKET_SIZE
    title(s, c['title'], 0.6, 0.38, 8.0, 0.75, size=20)
    text(s, 0.6, 1.18, 8.8, 0.32, c['method'], size=9.5, color=T.BODY)
    w, gap = 2.05, 0.2
    for i, (stage, size, geo, value, unit, desc) in enumerate(c['stages']):
        x = 0.6 + i * (w + gap)
        top = 2.35 - i * 0.2
        last = i == len(c['stages']) - 1
        fill = [T.PURPLE, T.LAV2, T.WHITE, T.WHITE][i]
        fg = T.WHITE if i == 0 else T.INK
        rect(s, x, top, w, 4.62 - top, fill=fill, line=T.BORDER if i >= 2 else None, radius=0.16, lw=0.75)
        _, cw = chip(s, stage, x + 0.16, top + 0.14, bg=T.WHITE if i == 0 else T.LAV2, fg=T.DEEP, size=7.5,
                     h=0.26)
        if size:
            chip(s, size, x + 0.22 + cw, top + 0.14, bg=ITEC_BLUE, fg=T.WHITE, size=7.5, h=0.26)
        text(s, x + 0.16, top + 0.48, w - 0.3, 0.24, geo, size=10, font=T.LABEL_FONT, color=fg)
        text(s, x + 0.16, top + 0.74, w - 0.3, 0.46, value, size=24, font=T.TITLE_FONT,
             color=T.WHITE if i == 0 else (T.PURPLE if last else T.DEEP))
        text(s, x + 0.16, top + 1.22, w - 0.3, 0.2, unit, size=8, font=T.LABEL_FONT,
             color=T.LAV2 if i == 0 else T.BODY)
        text(s, x + 0.16, top + 1.46, w - 0.3, 0.6, desc, size=8.5, color=T.WHITE if i == 0 else T.BODY,
             spacing=1.08)
        if i < 3:
            rect(s, x + w + 0.04, 3.5, 0.12, 0.26, fill=T.PURPLE, shape=MSO_SHAPE.CHEVRON)
    text(s, 0.6, 4.7, 8.8, 0.3, c['footer'], size=9, font=T.LABEL_FONT, color=T.DEEP, spacing=1.0)
    text(s, 0.6, 5.02, 8.8, 0.22, c['source'], size=6.5, color=T.BODY)
    notes(s, c['notes'])


def s_competition(s):
    c = K.SLIDE_COMPETITION
    I3.s_competition(s)
    lx, lw, cw = 0.6, 2.35, 1.33   # slides_dsc.s_competition table geometry
    for label, a, b in c['groups']:
        x = lx + lw + a * cw + 0.06
        rect(s, x, 1.15, (b - a + 1) * cw - 0.12, 0.17, fill=ITEC_BLUE if label == 'DIRECT' else T.GRAY,
             radius=0.085)
        text(s, x, 1.15, (b - a + 1) * cw - 0.12, 0.17, label, size=6.5, font=T.LABEL_FONT, color=T.WHITE,
             align=C, anchor=MID)


def s_uvp(s):
    c = K.SLIDE_UVP
    title(s, c['title'], 0.6, 0.45, 6.0, 0.5, size=26)
    text(s, 0.6, 1.02, 8.6, 0.5, c['lead'], size=12, color=T.BODY, spacing=1.1)
    # Venn: own creation (top-left), web access (top-right), shared worlds (bottom)
    d = 2.05
    spots = [(0.75, 1.6), (2.15, 1.6), (1.45, 2.85)]
    for (x, y), fill in zip(spots, [T.PURPLE, ITEC_BLUE, T.BLOB_MAIN]):
        _alpha(rect(s, x, y, d, d, fill=fill, line=fill, shape=MSO_SHAPE.OVAL, lw=1.25), 22000)
    boxes = [(0.92, 2.1, 1.18, 'l'), (2.92, 2.1, 1.18, 'r'), (1.8, 4.0, 1.35, 'c')]   # inside each circle only
    for (head, desc), (x, y, bw, a) in zip(c['circles'], boxes):
        al = {'l': L, 'r': R, 'c': C}[a]
        text(s, x, y, bw, 0.26, head, size=11, font=T.TITLE_FONT, color=T.DEEP, align=al)
        text(s, x, y + 0.27, bw, 0.6, desc, size=7.5, color=T.INK, align=al, spacing=1.0)
    chip(s, 'SCANVAS', 2.48, 3.18, bg=T.PURPLE, fg=T.WHITE, size=9, h=0.32, align='c')
    # mini matrix: which compared product combines all three
    players = ['Faber-Castell ARt Alive', 'Quiver', 'WAGZAK JUMP', 'teamLab Sketch Aquarium', 'Scanvas']
    marks = [(1, 0, 0), (1, 0, 0), (1, 0, 0), (1, 0, 1), (1, 1, 1)]
    x0, y0, cw, rh = 5.0, 1.75, 0.95, 0.4
    card(s, x0, y0 - 0.1, 4.4, 0.5 + rh * len(players) + 0.12, radius=0.14)
    for j, (head, _) in enumerate(c['circles']):
        text(s, x0 + 1.55 + j * cw, y0, cw, 0.36, head, size=8, font=T.LABEL_FONT, color=T.PURPLE, align=C,
             anchor=MID)
    for i, (p, m) in enumerate(zip(players, marks)):
        y = y0 + 0.42 + i * rh
        own = p == 'Scanvas'
        if own:
            rect(s, x0 + 0.08, y + 0.02, 4.24, rh - 0.04, fill=T.LAV2, radius=0.1)
        else:
            rect(s, x0 + 0.15, y, 4.1, 0.01, fill=T.BORDER)
        text(s, x0 + 0.18, y, 1.45, rh, p, size=8.5, font=T.TITLE_FONT if own else T.LABEL_FONT,
             color=T.DEEP if own else T.INK, anchor=MID, spacing=0.95)
        for j, v in enumerate(m):
            cx = x0 + 1.55 + j * cw + cw / 2
            if v:
                icon(s, 'check', 'purple' if own else 'green', cx - 0.1, y + rh / 2 - 0.1, 0.2)
            else:
                text(s, cx - 0.3, y, 0.6, rh, '—', size=9, color=T.GRAY, align=C, anchor=MID)
    text(s, 5.0, 4.62, 4.4, 0.55, c['proof'], size=8, color=T.BODY, spacing=1.05)
    notes(s, c['notes'])


def s_team(s):
    c = K.SLIDE_TEAM
    title(s, c['title'], 0.6, 0.45, 8.0, 0.5, size=24)
    w, top, h = 4.3, 1.25, 2.95
    for i, (name, role, photo, focus) in enumerate(c['team']):
        x = 0.6 + i * (w + 0.2)
        card(s, x, top, w, h, radius=0.18)
        pw = Image.open(T.GEN / photo).size
        ph = 2.35
        pic(s, photo, x + 0.25, top + 0.3, w=ph * pw[0] / pw[1], h=ph, alt=name)
        tx = x + 2.5
        text(s, tx, top + 0.32, w - 2.65, 0.62, name, size=15, font=T.TITLE_FONT, spacing=0.95)
        chip(s, role, tx, top + 1.0, bg=T.LAV2, fg=T.DEEP, size=8, h=0.28)
        for j, f in enumerate(focus):
            y = top + 1.5 + j * 0.32
            icon(s, 'check', 'purple', tx, y + 0.06, 0.16)
            text(s, tx + 0.24, y, w - 2.85, 0.28, f, size=9.5, anchor=MID)
    rect(s, 0.6, 4.4, 8.8, 0.75, fill=T.LAV2, radius=0.16)
    icon(s, 'quote', 'deep', 0.85, 4.6, 0.32)
    text(s, 1.35, 4.4, 7.8, 0.75, c['closing'], size=13, font=T.TITLE_FONT, color=T.INK, anchor=MID)
    notes(s, c['notes'])


def s_ask(s):
    c = K.SLIDE_ASK
    title(s, c['title'], 0.6, 0.45, 8.0, 0.5, size=26)
    text(s, 0.6, 1.02, 8.8, 0.3, c['lead'], size=12, color=T.BODY)
    w, h = 4.3, 1.68
    for i, (ic, head, detail, stage) in enumerate(c['asks']):
        x, y = 0.6 + (i % 2) * (w + 0.2), 1.55 + (i // 2) * (h + 0.18)
        hi = i == 0
        rect(s, x, y, w, h, fill=T.PURPLE if hi else T.WHITE, line=None if hi else T.BORDER, radius=0.18, lw=0.75)
        icon_circle(s, ic, x + 0.22, y + 0.22, 0.56, bg=T.WHITE if hi else T.LAV2)
        chip(s, 'UNLOCKS · ' + stage, x + w - 0.2, y + 0.24, bg=T.WHITE if hi else T.LAV2, fg=T.DEEP, size=7,
             h=0.24, align='r')
        text(s, x + 0.22, y + 0.9, w - 0.44, 0.32, head, size=14, font=T.TITLE_FONT,
             color=T.WHITE if hi else T.INK)
        text(s, x + 0.22, y + 1.22, w - 0.44, 0.42, detail, size=9.5, color=T.LAV2 if hi else T.BODY,
             spacing=1.05)
    notes(s, c['notes'])


def s_contact(s):
    c = K.SLIDE_CONTACT
    title(s, c['title'], 0.7, 0.85, 4.5, 0.7, size=40)
    rect(s, 0.72, 1.62, 3.4, 0.03, fill=ITEC_BLUE)
    for i, (ic, lbl, val, url) in enumerate(c['rows']):
        y = 2.0 + i * 0.66
        icon_circle(s, ic, 0.72, y, 0.48)
        text(s, 1.35, y - 0.02, 3.5, 0.2, lbl, size=8, font=T.LABEL_FONT, color=T.PURPLE)
        box = text(s, 1.35, y + 0.18, 3.5, 0.32, val, size=15, font=T.TITLE_FONT, anchor=MID)
        box.text_frame.paragraphs[0].runs[0].hyperlink.address = url
    text(s, 0.72, 3.5, 4.5, 0.26, c['people'], size=10, font=T.LABEL_FONT, color=T.INK, anchor=MID)
    text(s, 0.72, 3.8, 4.5, 0.24, K.TRACK, size=8.5, color=T.BODY, anchor=MID)
    card(s, 5.5, 0.75, 3.9, 4.35, radius=0.2)
    lw, lh = Image.open(T.LOGO).size
    ph = 0.7
    pic(s, T.LOGO, 5.5 + (3.9 - ph * lw / lh) / 2, 1.05, h=ph, alt='Scanvas logo')
    text(s, 5.5, 1.85, 3.9, 0.45, 'SCANVAS', size=24, font=T.TITLE_FONT, align=C, anchor=MID)
    qr, _, _ = pic(s, T.ASSETS / c['qr'], 6.55, 2.45, w=1.8, h=1.8, alt='QR code to the Scanvas digital trial')
    qr.click_action.hyperlink.address = K.SLIDE_4_SOLUTION['qr_url']
    chip(s, c['qr_label'], 7.45, 4.42, bg=T.PURPLE, fg=T.WHITE, size=8, h=0.3, align='c')
    notes(s, c['notes'])


def s_divider(s):
    text(s, 0.6, 2.2, 8.8, 0.8, K.SUPPORTING, size=40, font=T.TITLE_FONT, align=C, anchor=MID)
    rect(s, 4.0, 3.1, 2.0, 0.04, fill=ITEC_BLUE)
    text(s, 0.6, 3.3, 8.8, 0.3, getattr(K, 'SUPPORTING_SUB', 'Appendix · market data · technology & child data'),
         size=11, color=T.BODY,
         align=C, anchor=MID)


def s_funds(s):
    c = K.SLIDE_FUNDS
    title(s, c['title'], 0.6, 0.42, 7.6, 0.8, size=20)
    # capital need
    rect(s, 0.6, 1.35, 2.75, 3.85, fill=T.PURPLE, radius=0.18)
    text(s, 0.8, 1.55, 2.4, 0.2, c['label'], size=7.5, font=T.LABEL_FONT, color=T.LAV2)
    text(s, 0.8, 1.78, 2.4, 0.62, c['amount'], size=36, font=T.TITLE_FONT, color=T.WHITE, anchor=MID)
    text(s, 0.8, 2.42, 2.4, 0.2, c['usd'], size=7.5, color=T.LAV2)
    rect(s, 0.8, 2.75, 2.35, 0.01, fill=T.LAV2)
    text(s, 0.8, 2.9, 2.35, 1.3, c['purpose'], size=10, color=T.WHITE, spacing=1.12)
    chip(s, '12 MONTHS', 0.8, 4.68, bg=T.WHITE, fg=T.DEEP, size=8, h=0.3, ic=('calendar-check', 'purple'))
    # allocation bars (width = share of Rp200M)
    x0, w = 3.55, 3.15
    total = sum(a[3] for a in c['alloc'])
    text(s, x0, 1.35, w, 0.22, c['alloc_title'], size=8, font=T.LABEL_FONT, color=T.PURPLE)
    for i, (ic, use, detail, amt) in enumerate(c['alloc']):
        y = 1.68 + i * 0.52
        icon_circle(s, ic, x0, y, 0.34)
        text(s, x0 + 0.44, y - 0.02, 1.8, 0.22, use, size=9.5, font=T.LABEL_FONT, anchor=MID)
        text(s, x0 + 1.9, y - 0.02, w - 1.9, 0.22, f'Rp{amt}M · {round(100 * amt / total)}%', size=9.5,
             font=T.TITLE_FONT, color=T.DEEP, align=R, anchor=MID)
        bw = w - 0.44
        rect(s, x0 + 0.44, y + 0.24, bw, 0.08, fill=T.LAV2, radius=0.04)
        rect(s, x0 + 0.44, y + 0.24, max(0.1, bw * amt / total), 0.08, fill=T.PURPLE, radius=0.04)
        text(s, x0 + 0.44, y + 0.33, bw, 0.16, detail, size=7, color=T.BODY)
    # milestones
    card(s, 6.9, 1.35, 2.5, 3.45, radius=0.16)
    text(s, 7.05, 1.47, 2.25, 0.36, c['unlocks_title'], size=7.5, font=T.LABEL_FONT, color=T.PURPLE, spacing=1.0)
    for i, u in enumerate(c['unlocks']):
        y = 1.9 + i * 0.4
        icon(s, 'check', 'purple', 7.05, y + 0.06, 0.16)
        text(s, 7.28, y, 2.05, 0.3, u, size=8.5, anchor=MID, spacing=1.0)
    text(s, 3.55, 4.88, 5.85, 0.34, c['basis'], size=7, color=T.BODY, spacing=1.05)
    notes(s, c['notes'])


def s_exit(s):
    c = K.SLIDE_EXIT
    title(s, c['title'], 0.6, 0.42, 7.6, 0.8, size=20)
    text(s, 0.6, 1.2, 8.8, 0.26, c['lead'], size=10.5, color=T.BODY)
    # indicative timeline
    n = len(c['timeline'])
    step = 8.8 / n
    rect(s, 0.75, 1.76, step * (n - 1), 0.03, fill=ITEC_BLUE)
    for i, (period, head, detail) in enumerate(c['timeline']):
        x = 0.6 + i * step
        last = i == n - 1
        d = 0.22 if last else 0.16
        oval(s, x + 0.15 - d / 2 + 0.0, 1.775 - d / 2, d, fill=T.PURPLE if last else ITEC_BLUE)
        text(s, x, 1.95, step - 0.15, 0.18, period, size=7.5, font=T.LABEL_FONT, color=T.PURPLE)
        text(s, x, 2.13, step - 0.15, 0.26, head, size=11.5, font=T.TITLE_FONT)
        text(s, x, 2.4, step - 0.2, 0.36, detail, size=8, color=T.BODY, spacing=1.0)
    # exit options
    w, gap, top, h = 2.08, 0.16, 2.95, 1.78
    for i, (tag, head, detail, primary) in enumerate(c['options']):
        x = 0.6 + i * (w + gap)
        rect(s, x, top, w, h, fill=T.PURPLE if primary else T.WHITE, line=None if primary else T.BORDER,
             radius=0.16, lw=0.75)
        chip(s, tag, x + 0.16, top + 0.15, bg=T.WHITE if primary else T.LAV2, fg=T.DEEP, size=7, h=0.24)
        text(s, x + 0.16, top + 0.5, w - 0.32, 0.3, head, size=12, font=T.TITLE_FONT,
             color=T.WHITE if primary else T.INK)
        text(s, x + 0.16, top + 0.84, w - 0.32, 0.9, detail, size=8.5, color=T.LAV2 if primary else T.BODY,
             spacing=1.08)
    rect(s, 0.6, 4.86, 8.8, 0.36, fill=T.LAV2, radius=0.18)
    text(s, 0.75, 4.86, 8.5, 0.36, c['value'], size=9, font=T.LABEL_FONT, color=T.DEEP, align=C, anchor=MID)
    notes(s, c['notes'])
