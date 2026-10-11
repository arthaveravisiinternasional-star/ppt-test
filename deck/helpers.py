"""Thin python-pptx helpers. All coordinates are inches on the 10 x 5.625 in Freesia canvas."""
from pptx.util import Inches as _Inches, Pt as _Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from PIL import Image

from . import theme as T

L, C, R = PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.RIGHT
TOP, MID = MSO_ANCHOR.TOP, MSO_ANCHOR.MIDDLE

SCALE = 1.0   # build_deck_itec4.py sets 2.0: the official ITEC template is 20 x 11.25 in (2x this canvas)


def Inches(v):
    return _Inches(v * SCALE)


def Pt(v):
    return _Pt(v * SCALE)


def rgb(h):
    return RGBColor.from_string(h)


ARROW, ARROW_FONT = '→', 'Segoe UI'   # Fredoka / Varela Round have no arrow glyph


def _run(p, t, size, color, font, bold=False, italic=False):
    """Add styled run(s); a newline inside t becomes a soft line break, arrows use ARROW_FONT."""
    r = None
    for i, line in enumerate(t.split('\n')):
        if i:
            p.add_line_break()
        for j, part in enumerate(line.split(ARROW)):
            for txt, fnt in ([(ARROW, ARROW_FONT)] if j else []) + [(part, font)]:
                if not txt:
                    continue
                r = p.add_run()
                r.text = txt
                f = r.font
                f.size = Pt(size); f.name = T.FONT_ALIAS.get(fnt, fnt); f.bold = bold or font in T.BOLD_FONTS; f.italic = italic
                f.color.rgb = rgb(color)
    return r


def _frame(tf, anchor=TOP):
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor


def text(s, x, y, w, h, paras, size=12, color=T.INK, font=T.BODY_FONT, align=L, anchor=TOP,
         spacing=None, shape=None):
    """paras: str | [str, ...] | [[(text, {size,color,font}), ...], ...] (one inner list per paragraph)."""
    if shape is None:
        shape = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    _frame(tf, anchor)
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.line_spacing = spacing
        for t, o in (para if isinstance(para, list) else [(para, {})]):
            _run(p, t, o.get('size', size), o.get('color', color), o.get('font', font))
    return shape


def title(s, t, x, y, w, h, size=24, color=T.INK, align=L):
    """Fill the layout's own title placeholder (keeps slide outline / accessibility)."""
    ph = s.shapes.title
    if ph is None:
        return text(s, x, y, w, h, t, size=size, color=color, font=T.TITLE_FONT, align=align, spacing=0.95)
    ph.left, ph.top, ph.width, ph.height = Inches(x), Inches(y), Inches(w), Inches(h)
    tf = ph.text_frame
    tf.clear()
    _frame(tf)
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = 0.95
    _run(p, t, size, color, T.TITLE_FONT)
    return ph


def rect(s, x, y, w, h, fill=None, line=None, radius=None, shape=None, lw=1.0, rot=0):
    st = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    sh = s.shapes.add_shape(st, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.shadow.inherit = False
    if fill:
        sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = rgb(line); sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    if radius and st == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = min(0.5, radius / min(w, h))
    if rot:
        sh.rotation = rot
    return sh


def card(s, x, y, w, h, fill=T.WHITE, radius=0.16):
    """DESIGN.md card: white surface on canvas, one Garis Lembut hairline, no shadow."""
    return rect(s, x, y, w, h, fill=fill, line=T.BORDER, radius=radius, lw=0.75)


def oval(s, x, y, d, fill=None, line=None, lw=1.0):
    return rect(s, x, y, d, d, fill=fill, line=line, shape=MSO_SHAPE.OVAL, lw=lw)


def img_size(name, w=None, h=None):
    im = Image.open(T.GEN / name)
    ar = im.size[0] / im.size[1]
    return (w, w / ar) if w else (h * ar, h)


def pic(s, path, x, y, w=None, h=None, alt=None):
    """path: file name inside assets/gen, or an absolute Path."""
    path = str(path if '/' in str(path) or '\\' in str(path) else T.GEN / path)
    im = Image.open(path)
    ar = im.size[0] / im.size[1]
    if w and not h:
        h = w / ar
    elif h and not w:
        w = h * ar
    p = s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    if alt:
        p._element.nvPicPr.cNvPr.set('descr', alt)
    return p, w, h


def icon(s, name, color, x, y, d):
    """color: purple | white | ink | green | gray | yellow | deep"""
    return s.shapes.add_picture(str(T.ICONS / f'{name}_{color}.png'), Inches(x), Inches(y), Inches(d), Inches(d))


def icon_circle(s, name, x, y, d, bg=T.LAV2, ic='purple', ratio=0.52, line=None):
    oval(s, x, y, d, fill=bg, line=line)
    k = d * ratio
    icon(s, name, ic, x + (d - k) / 2, y + (d - k) / 2, k)


def num_circle(s, n, x, y, d, bg=T.PURPLE, fg=T.WHITE, size=12):
    c = oval(s, x, y, d, fill=bg)
    text(s, 0, 0, 0, 0, str(n), size=size, color=fg, font=T.TITLE_FONT, align=C, anchor=MID, shape=c)
    return c


def chip_width(t, size, h=0.28, has_icon=False):
    return len(t) * size * 0.68 / 72 + 0.36 + (h * 0.8 if has_icon else 0)


def chip(s, t, x, y, bg=T.PURPLE, fg=T.WHITE, size=8.5, h=0.28, w=None, line=None, ic=None, align='l'):
    """Pill label. ic=(icon_name, icon_color). align l/r/c anchors x."""
    w = w or chip_width(t, size, h, bool(ic))
    x = x - w if align == 'r' else x - w / 2 if align == 'c' else x
    c = rect(s, x, y, w, h, fill=bg, line=line, radius=h / 2, lw=1.25)
    tf = c.text_frame
    _frame(tf, MID)
    tf.margin_left = Inches(0.12 + (h * 0.78 if ic else 0)); tf.margin_right = Inches(0.1)
    p = tf.paragraphs[0]
    p.alignment = L if ic else C
    _run(p, t, size, fg, T.LABEL_FONT)
    if ic:
        k = h * 0.62
        icon(s, ic[0], ic[1], x + 0.12, y + (h - k) / 2, k)
    return c, w


def arrow_chips(s, items, x, y, h=0.3, size=8.5, gap=0.24, bg=T.LAV2, fg=T.DEEP, last=None, arrow='purple'):
    """Row of chips joined by arrow icons. last=(bg, fg) restyles the final chip."""
    cx = x
    for i, t in enumerate(items):
        b, f = last if (last and i == len(items) - 1) else (bg, fg)
        _, w = chip(s, t, cx, y, bg=b, fg=f, size=size, h=h)
        cx += w
        if i < len(items) - 1:
            k = 0.15
            icon(s, 'arrow-right', arrow, cx + (gap - k) / 2, y + (h - k) / 2, k)
            cx += gap
    return cx


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t
