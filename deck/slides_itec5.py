"""Slides for the ITEC SEA 2026 7-minute pitch edition (v5): 12 main slides + appendix on the committee template.
Same grid and helpers as slides_itec4 (helpers.SCALE = 2). Keep content between y 0.3 and 5.3 (template bars).
Imported after build_deck_itec5.py swaps deck.content for deck.content_itec5."""
import math

from PIL import Image
from pptx.enum.shapes import MSO_SHAPE

from . import content as K
from . import theme as T
from .helpers import (L, C, MID, R, Inches, text, title, rect, card, oval, pic, icon, num_circle, chip, img_size,
                      notes)
from .slides_itec3 import _flow
from .slides_itec4 import ITEC_BLUE


def _lines(t, size, w, k=0.43):
    """Rough line count for Roboto Condensed (average glyph ~0.43 em)."""
    return max(1, math.ceil(len(t) * size * k / 72 / w))


def _qr(s, x, y, d):
    qr, _, _ = pic(s, T.ASSETS / K.SLIDE_4_SOLUTION['qr'], x, y, w=d, h=d, alt='QR code to the Scanvas digital trial')
    qr.click_action.hyperlink.address = K.SLIDE_4_SOLUTION['qr_url']


# ---------------------------------------------------------------- 1
def s_cover(s):
    c = K.COVER5
    pic(s, T.LOGO, 0.72, 0.6, h=0.5, alt='Scanvas logo')
    text(s, 0.72, 1.24, 4.0, 0.24, c['label'], size=8.5, font=T.LABEL_FONT, color=ITEC_BLUE, anchor=MID)
    title(s, c['title'], 0.7, 1.5, 4.2, 0.75, size=46)
    text(s, 0.72, 2.27, 3.9, 0.7, c['subtitle'], size=18, color=T.PURPLE, font=T.TITLE_FONT, spacing=0.95)
    text(s, 0.72, 2.95, 3.75, 0.75, c['summary'], size=12, color=T.BODY, spacing=1.08)
    y = 3.86
    for name, role in c['presenters']:
        text(s, 0.72, y, 3.9, 0.24, [[(name, {'font': T.LABEL_FONT, 'size': 10.5}),
                                      ('  —  ' + role, {'size': 9, 'color': T.BODY})]], anchor=MID)
        y += 0.27
    text(s, 0.72, 4.48, 3.9, 0.22, c['track'], size=9, font=T.LABEL_FONT, color=ITEC_BLUE, anchor=MID)
    box = text(s, 0.72, 4.74, 2.0, 0.26, c['website'], size=12, font=T.TITLE_FONT, color=T.PURPLE, anchor=MID)
    box.text_frame.paragraphs[0].runs[0].hyperlink.address = 'https://www.scanvas.id/'
    pic(s, 'paper_whale_tilt.png', 4.55, 1.05, w=2.55, alt='A child-colored Scanvas whale sheet')
    rect(s, 7.08, 3.15, 0.16, 0.34, fill=T.PURPLE, shape=MSO_SHAPE.CHEVRON)
    oval(s, 7.25, 2.75, 2.2, fill=T.LAV2)
    pic(s, 'ph_world.png', 7.62, 2.45, h=2.85, alt='The same whale swimming in the Scanvas ocean world')
    notes(s, c['notes'])


# ---------------------------------------------------------------- 2
def s_problem(s):
    c = K.PROBLEM5
    title(s, c['title'], 0.6, 0.45, 8.0, 0.5, size=21)
    y = 1.2
    for i, (lbl, t) in enumerate(c['problems']):
        num_circle(s, i + 1, 0.6, y + 0.02, 0.32, size=10)
        text(s, 1.05, y, 3.8, 0.2, lbl, size=7.5, font=T.LABEL_FONT, color=T.PURPLE)
        text(s, 1.05, y + 0.2, 3.85, 0.5, t, size=10.5, color=T.INK, spacing=1.08)
        y += 0.78
    rect(s, 0.6, 3.6, 4.35, 0.98, fill=T.PURPLE, radius=0.16)
    icon(s, 'sparkles', 'white', 0.8, 3.76, 0.28)
    text(s, 1.2, 3.6, 3.6, 0.98, c['question'], size=13, font=T.TITLE_FONT, color=T.WHITE, anchor=MID, spacing=1.0)
    icon(s, 'flask-conical', 'purple', 0.6, 4.74, 0.2)
    text(s, 0.86, 4.68, 4.1, 0.34, c['tested'], size=9, font=T.LABEL_FONT, color=T.DEEP, anchor=MID, spacing=1.0)
    pic(s, c['photo'], 5.2, 1.25, w=4.2, h=4.2 * 9 / 16, alt='A child coloring a Scanvas fish sheet in the pilot')
    for i, a in enumerate(K.PILOT_ART):
        pic(s, a, 5.2 + i * 1.08, 3.7, w=0.96, h=0.96, alt='Artwork colored by a child in the pilot')
    text(s, 5.2, 4.72, 4.2, 0.22, c['caption'], size=8, font=T.LABEL_FONT, color=T.PURPLE, anchor=MID)
    notes(s, c['notes'])


# ---------------------------------------------------------------- 3
def s_solution(s):
    c = K.SOLUTION5
    title(s, c['title'], 0.6, 0.45, 8.0, 0.5, size=26)
    tw, gap, ty, th = 1.98, 0.29, 1.35, 2.05
    for i, (lbl, cap, img, fit) in enumerate(c['steps']):
        x = 0.6 + i * (tw + gap)
        card(s, x, ty, tw, th)
        w, h = img_size(img, w=tw - 0.06) if fit == 'w' else img_size(img, h=th - 0.1)
        if h > th - 0.1:
            w, h = img_size(img, h=th - 0.1)
        pic(s, img, x + (tw - w) / 2, ty + (th - h) / 2, w=w, h=h, alt=f'{lbl} step')
        num_circle(s, i + 1, x - 0.08, ty - 0.12, 0.36, size=11)
        text(s, x, ty + th + 0.14, tw, 0.3, lbl, size=15, font=T.TITLE_FONT, color=T.DEEP)
        text(s, x, ty + th + 0.46, tw, 0.6, cap, size=10.5, color=T.BODY, spacing=1.08)
        if i < len(c['steps']) - 1:
            rect(s, x + tw + 0.09, ty + th / 2 - 0.14, 0.12, 0.28, fill=T.PURPLE, shape=MSO_SHAPE.CHEVRON)
    rect(s, 0.6, 4.66, 8.8, 0.48, fill=T.LAV2, radius=0.24)
    text(s, 0.85, 4.66, 5.9, 0.48, c['statement'], size=12.5, font=T.TITLE_FONT, color=T.DEEP, anchor=MID)
    chip(s, c['chip'], 9.25, 4.76, bg=T.PURPLE, fg=T.WHITE, size=8.5, h=0.28, ic=('check', 'white'), align='r')
    notes(s, c['notes'])


# ---------------------------------------------------------------- 4
def s_demo(s):
    c = K.DEMO5
    chip(s, c['tag'], 0.6, 0.5, bg=T.MINT, fg=T.GREEN, size=8, h=0.28, ic=('play', 'green'))
    title(s, c['title'], 0.6, 0.9, 2.6, 0.85, size=24)
    text(s, 0.6, 1.82, 2.55, 0.8, c['body'], size=10, color=T.BODY, spacing=1.1)
    text(s, 0.6, 2.62, 2.5, 0.2, c['live_title'], size=7.5, font=T.LABEL_FONT, color=T.PURPLE)
    for i, it in enumerate(c['live']):
        y = 2.86 + i * 0.32
        icon(s, 'check', 'purple', 0.6, y + 0.05, 0.16)
        text(s, 0.84, y, 2.4, 0.26, it, size=10.5, anchor=MID)
    x, y, w = 3.3, 0.6, 6.1
    h = w * 9 / 16
    rect(s, x - 0.05, y - 0.05, w + 0.1, h + 0.1, fill=T.WHITE, line=T.BORDER, radius=0.06, lw=0.75)
    movie = s.shapes.add_movie(str(T.ASSETS / c['video']), Inches(x), Inches(y), Inches(w), Inches(h),
                               poster_frame_image=str(T.ASSETS / c['poster']), mime_type='video/mp4')
    movie.name = 'Scanvas digital trial demo'
    rect(s, x + 0.9, y + h + 0.16, w - 1.8, 0.4, fill=T.LAV2, radius=0.2)
    text(s, x + 0.9, y + h + 0.16, w - 1.8, 0.4, c['caption'], size=11, font=T.TITLE_FONT, color=T.DEEP, align=C,
         anchor=MID)
    notes(s, c['notes'])


# ---------------------------------------------------------------- 5
def s_traction(s):
    c = K.TRACTION5
    title(s, c['title'], 0.6, 0.45, 8.0, 0.5, size=24)
    n = len(c['metrics'])
    mw = (4.5 - 0.1 * (n - 1)) / n
    for i, (v, lbl) in enumerate(c['metrics']):
        x = 0.6 + i * (mw + 0.1)
        card(s, x, 1.2, mw, 1.05, radius=0.14)
        text(s, x, 1.24, mw, 0.48, v, size=26, font=T.TITLE_FONT, color=T.DEEP, align=C, anchor=MID)
        text(s, x + 0.06, 1.72, mw - 0.12, 0.5, lbl, size=8, font=T.LABEL_FONT, align=C, spacing=1.0)
    card(s, 0.6, 2.38, 4.5, 1.4, radius=0.14)
    for j, (head, items, done) in enumerate([('SHIPPED', c['shipped'], True), ('BUILDING', c['building'], False)]):
        x = 0.78 + j * 2.2
        chip(s, head, x, 2.51, bg=T.MINT if done else T.LAV2, fg=T.GREEN if done else T.DEEP, size=7.5, h=0.26)
        for k, it in enumerate(items):
            y = 2.89 + k * 0.28
            icon(s, 'check' if done else 'arrow-right', 'green' if done else 'purple', x, y + 0.05, 0.15)
            text(s, x + 0.22, y, 1.95, 0.26, it, size=9.5, anchor=MID)
    for j, (lbl, t) in enumerate([(c['proved_label'], c['proved']), (c['next_label'], c['next'])]):
        x = 0.6 + j * 2.3
        rect(s, x, 3.9, 2.2, 0.9, fill=T.LAV if j == 0 else T.PALEY, radius=0.12)
        text(s, x + 0.14, 3.96, 1.95, 0.2, lbl, size=7, font=T.LABEL_FONT, color=T.PURPLE if j == 0 else T.AMBER_DARK)
        text(s, x + 0.14, 4.18, 1.95, 0.6, t, size=10, font=T.LABEL_FONT, color=T.INK, spacing=1.05)
    chip(s, c['stage'], 0.6, 4.92, bg=T.WHITE, fg=T.DEEP, size=7.5, h=0.24, line=T.DEEP)
    big, *small = c['photos']
    pic(s, big, 5.3, 1.2, w=4.1, h=4.1 * 9 / 16, alt='Scanvas demo with children during the TPA pilot')
    sw = (4.1 - 0.1) / 2
    for i, p in enumerate(small):
        pic(s, p, 5.3 + i * (sw + 0.1), 3.6, w=sw, h=sw * 9 / 16, alt='Children coloring Scanvas sheets in the pilot')
    text(s, 5.3, 4.82, 4.1, 0.22, c['caption'], size=8, font=T.LABEL_FONT, color=T.PURPLE, anchor=MID)
    notes(s, c['notes'])


# ---------------------------------------------------------------- 6
def s_market(s):
    c = K.MARKET5
    title(s, c['title'], 0.6, 0.45, 8.0, 0.5, size=24)
    widths, top, h, gap = [5.3, 4.5, 3.7], 1.15, 0.9, 0.1   # funnel: TAM (widest) -> SAM -> beachhead
    for i, ((lbl, v, unit, place), w) in enumerate(zip(c['tiers'], widths)):
        x = 0.6 + (5.3 - w) / 2
        y = top + i * (h + gap)
        last = i == len(widths) - 1
        rect(s, x, y, w, h, fill=T.PURPLE if last else (T.LAV2 if i == 1 else T.WHITE),
             line=None if i else T.BORDER, radius=0.16, lw=0.75)
        fg = T.WHITE if last else T.INK
        chip(s, lbl, x + 0.18, y + 0.14, bg=T.WHITE if last else ITEC_BLUE, fg=T.DEEP if last else T.WHITE, size=7.5,
             h=0.26)
        text(s, x + 0.18, y + 0.44, 1.7, 0.4, v, size=22, font=T.TITLE_FONT, color=T.WHITE if last else T.DEEP,
             anchor=MID)
        text(s, x + 1.85, y + 0.14, w - 2.0, 0.7, [[(place, {'font': T.LABEL_FONT, 'size': 10.5})], unit], size=9,
             color=fg, anchor=MID, spacing=1.05)
    icon(s, 'map-pin', 'purple', 0.6, 4.17, 0.2)
    text(s, 0.86, 4.1, 5.0, 0.34, c['why'], size=10, font=T.LABEL_FONT, color=T.DEEP, anchor=MID)
    for i, (lbl, v) in enumerate([c['buyer'], c['user']]):
        y = 1.15 + i * 0.95
        card(s, 6.1, y, 3.3, 0.85, radius=0.14)
        text(s, 6.28, y + 0.1, 3.0, 0.2, lbl, size=7.5, font=T.LABEL_FONT, color=T.PURPLE)
        text(s, 6.28, y + 0.32, 3.0, 0.48, v, size=11.5, font=T.TITLE_FONT, spacing=1.0)
    rect(s, 6.1, 3.05, 3.3, 1.39, fill=T.PALEY, radius=0.16)
    text(s, 6.28, 3.16, 3.0, 0.2, c['target_label'], size=7.5, font=T.LABEL_FONT, color=T.AMBER_DARK)
    text(s, 6.28, 3.38, 3.0, 0.5, c['target'], size=22, font=T.TITLE_FONT, color=T.INK, anchor=MID)
    text(s, 6.28, 3.92, 3.0, 0.45, c['target_note'], size=9, color=T.AMBER_DARK, spacing=1.0)
    text(s, 0.6, 4.62, 8.8, 0.4, c['source'], size=7.5, color=T.BODY, spacing=1.05)
    notes(s, c['notes'])


# ---------------------------------------------------------------- 7
def s_competition(s):
    c = K.COMPETITION5
    title(s, c['title'], 0.6, 0.45, 8.0, 0.5, size=24)
    text(s, 0.6, 1.15, 3.0, 0.45, c['intro'], size=10, color=T.BODY, spacing=1.05)
    for i, (name, kind) in enumerate(c['alternatives']):
        y = 1.72 + i * 0.6
        card(s, 0.6, y, 3.0, 0.5, radius=0.12)
        text(s, 0.75, y, 1.9, 0.5, name, size=10, font=T.LABEL_FONT, anchor=MID, spacing=1.0)
        chip(s, kind, 3.48, y + 0.13, bg=T.SLATE2, fg=T.INK, size=7, h=0.24, align='r')
    oval(s, 4.05, 1.55, 2.3, fill=T.LAV2)
    pic(s, 'ph_world.png', 4.5, 1.15, h=3.2, alt='A child’s whale in the shared Scanvas ocean, in the browser')
    for (lbl, (x, y)) in zip(c['callouts'], [(3.85, 1.25), (4.95, 2.55), (3.85, 3.85)]):
        chip(s, lbl, x, y, bg=T.PURPLE, fg=T.WHITE, size=7.5, h=0.26)
    card(s, 6.75, 1.15, 2.65, 3.3, radius=0.16)
    text(s, 6.92, 1.27, 2.4, 0.2, c['combines_title'], size=7.5, font=T.LABEL_FONT, color=T.PURPLE)
    for i, it in enumerate(c['combines']):
        y = 1.55 + i * 0.57
        icon(s, 'check', 'purple', 6.92, y + 0.04, 0.18)
        text(s, 7.16, y, 2.15, 0.52, it, size=10, spacing=1.05)
    rect(s, 0.6, 4.6, 8.8, 0.52, fill=T.PALEY, radius=0.26)
    text(s, 0.8, 4.6, 8.4, 0.52, c['next_test'], size=11.5, font=T.TITLE_FONT, color=T.AMBER_DARK, align=C,
         anchor=MID)
    notes(s, c['notes'])


# ---------------------------------------------------------------- 8
def s_business(s):
    c = K.BUSINESS5
    title(s, c['title'], 0.6, 0.45, 8.0, 0.5, size=24)
    w, top, h = 4.3, 1.15, 3.4
    for i, (tag, style, head, sub, price, flow, validate) in enumerate(c['paths']):
        x = 0.6 + i * (w + 0.2)
        primary = style == 'primary'
        rect(s, x, top, w, h, fill=T.WHITE, line=T.PURPLE if primary else T.BORDER, radius=0.18,
             lw=1.5 if primary else 0.75)
        chip(s, tag, x + 0.2, top + 0.18, bg=T.PURPLE if primary else T.MINT, fg=T.WHITE if primary else T.GREEN,
             size=7.5, h=0.26)
        text(s, x + 0.2, top + 0.56, w - 0.4, 0.34, head, size=15, font=T.TITLE_FONT)
        text(s, x + 0.2, top + 0.9, w - 0.4, 0.24, sub, size=10, color=T.BODY)
        rect(s, x + 0.15, top + 1.22, w - 0.3, 0.62, fill=T.LAV, radius=0.12)
        text(s, x + 0.3, top + 1.26, w - 0.6, 0.18, 'TEST PRICE', size=7, font=T.LABEL_FONT, color=T.PURPLE)
        text(s, x + 0.3, top + 1.44, w - 0.6, 0.36, price, size=17, font=T.TITLE_FONT, color=T.DEEP, anchor=MID)
        text(s, x + 0.2, top + 1.98, w - 0.4, 0.18, 'ACQUISITION', size=7, font=T.LABEL_FONT, color=T.PURPLE)
        _flow(s, ' → '.join(f.upper() for f in flow), x + 0.15, top + 2.18, w - 0.3, h=0.3, size=7.5)
        text(s, x + 0.2, top + 2.6, w - 0.4, 0.18, 'VALIDATE', size=7, font=T.LABEL_FONT, color=T.PURPLE)
        text(s, x + 0.2, top + 2.8, w - 0.4, 0.5, validate, size=10.5, font=T.LABEL_FONT, spacing=1.05)
    rect(s, 0.6, 4.68, 8.8, 0.44, fill=T.PALEY, radius=0.22)
    icon(s, 'flask-conical', 'yellow', 1.6, 4.79, 0.22)
    text(s, 1.9, 4.68, 7.2, 0.44, c['footer'], size=11, font=T.TITLE_FONT, color=T.AMBER_DARK, anchor=MID)
    notes(s, c['notes'])


# ---------------------------------------------------------------- 9
def s_team(s):
    c = K.TEAM5
    title(s, c['title'], 0.6, 0.45, 6.8, 0.5, size=22)
    chip(s, c['equity'], 9.4, 0.6, bg=T.LAV2, fg=T.DEEP, size=8, h=0.28, w=2.05, ic=('users', 'purple'), align='r')
    w, top, h = 4.3, 1.15, 2.72
    for i, (name, role, photo, focus) in enumerate(c['team']):
        x = 0.6 + i * (w + 0.2)
        card(s, x, top, w, h, radius=0.18)
        pw = Image.open(T.GEN / photo).size
        ph = 1.95
        pic(s, photo, x + 0.22, top + 0.38, w=ph * pw[0] / pw[1], h=ph, alt=name)
        tx = x + 2.3
        text(s, tx, top + 0.18, w - 2.4, 0.56, name, size=14, font=T.TITLE_FONT, spacing=0.95)
        chip(s, role, tx, top + 0.78, bg=T.LAV2, fg=T.DEEP, size=7.5, h=0.26)
        for j, f in enumerate(focus):
            y = top + 1.14 + j * 0.29
            icon(s, 'check', 'purple', tx, y + 0.06, 0.15)
            text(s, tx + 0.22, y, w - 2.55, 0.27, f, size=9.5, anchor=MID)
    for j, (lbl, t, fill) in enumerate([(c['execution_title'], c['execution'], T.WHITE),
                                        (c['next_title'], c['next'], T.LAV2)]):
        x = 0.6 + j * 4.5
        rect(s, x, 4.02, 4.3, 1.12, fill=fill, line=T.BORDER if j == 0 else None, radius=0.16, lw=0.75)
        text(s, x + 0.2, 4.12, 3.9, 0.2, lbl, size=7.5, font=T.LABEL_FONT, color=T.PURPLE)
        text(s, x + 0.2, 4.34, 3.9, 0.72, t, size=13, font=T.TITLE_FONT, anchor=MID, spacing=1.0)
    notes(s, c['notes'])


# ---------------------------------------------------------------- 10
def s_roadmap(s):
    c = K.ROADMAP5
    title(s, c['title'], 0.6, 0.45, 8.0, 0.5, size=22)
    w, gap, top, h = 2.08, 0.16, 1.15, 3.4
    for i, (tag, head, items, done) in enumerate(c['stages']):
        x = 0.6 + i * (w + gap)
        fill = T.PURPLE if done else (T.LAV2 if i == 1 else T.WHITE)
        fg = T.WHITE if done else T.INK
        rect(s, x, top, w, h, fill=fill, line=None if done or i == 1 else T.BORDER, radius=0.16, lw=0.75)
        chip(s, tag, x + 0.15, top + 0.15, bg=T.WHITE if done or i == 1 else T.LAV2, fg=T.DEEP, size=7.5, h=0.26)
        text(s, x + 0.15, top + 0.5, w - 0.3, 0.32, head, size=13, font=T.TITLE_FONT, color=fg)
        y = top + 0.95
        for it in items:
            n = _lines(it, 9.5, w - 0.55)
            icon(s, 'check' if done else 'arrow-right', 'white' if done else 'purple', x + 0.15, y + 0.05, 0.15)
            text(s, x + 0.38, y, w - 0.5, 0.2 * n + 0.04, it, size=9.5, color=fg, spacing=1.0)
            y += 0.2 * n + 0.2
        if i < 3:
            rect(s, x + w + 0.03, top + h / 2 - 0.13, 0.1, 0.26, fill=T.PURPLE, shape=MSO_SHAPE.CHEVRON)
    rect(s, 0.6, 4.68, 8.8, 0.44, fill=T.PURPLE, radius=0.22)
    text(s, 0.6, 4.68, 8.8, 0.44, c['footer'], size=12.5, font=T.TITLE_FONT, color=T.WHITE, align=C, anchor=MID)
    notes(s, c['notes'])


# ---------------------------------------------------------------- 11
def s_funds(s):
    c = K.FUNDS5
    title(s, c['title'], 0.6, 0.45, 8.0, 0.5, size=24)
    for i, (head, rows, margin) in enumerate(c['units']):     # unit economics
        y0 = 1.15 + i * 1.72
        card(s, 0.6, y0, 2.85, 1.62, radius=0.14)
        text(s, 0.76, y0 + 0.1, 2.6, 0.2, head, size=7.5, font=T.LABEL_FONT, color=T.PURPLE)
        for j, (lbl, v) in enumerate(rows):
            y = y0 + 0.36 + j * 0.3
            text(s, 0.76, y, 1.4, 0.28, lbl, size=10, color=T.BODY, anchor=MID)
            text(s, 2.0, y, 1.3, 0.28, v, size=12 if j == 2 else 11, font=T.TITLE_FONT,
                 color=T.DEEP if j == 2 else T.INK, align=R, anchor=MID)
        chip(s, margin, 0.76, y0 + 1.3, bg=T.MINT, fg=T.GREEN, size=7.5, h=0.24)
    card(s, 3.6, 1.15, 2.75, 3.34, radius=0.14)               # base case
    text(s, 3.78, 1.25, 2.4, 0.2, c['base_title'], size=7.5, font=T.LABEL_FONT, color=T.PURPLE)
    text(s, 3.78, 1.47, 2.45, 0.26, c['base_volume'], size=10, font=T.LABEL_FONT, anchor=MID)
    for j, (lbl, v) in enumerate(c['base_rows']):
        y = 1.86 + j * 0.62
        text(s, 3.78, y, 2.4, 0.2, lbl.upper(), size=7, font=T.LABEL_FONT, color=T.BODY)
        text(s, 3.78, y + 0.2, 2.4, 0.38, v, size=20, font=T.TITLE_FONT, color=T.DEEP if j != 1 else T.INK,
             anchor=MID)
    chip(s, c['base_margin'], 3.78, 3.82, bg=T.MINT, fg=T.GREEN, size=7.5, h=0.24)
    text(s, 3.78, 4.12, 2.45, 0.32, 'Founders’ model, Sep 2026 · assumptions', size=7.5, color=T.BODY)
    rect(s, 6.5, 1.15, 2.9, 3.34, fill=T.PURPLE, radius=0.16)  # capital required + use
    text(s, 6.68, 1.25, 2.6, 0.2, c['capital_title'], size=7.5, font=T.LABEL_FONT, color=T.LAV2)
    text(s, 6.68, 1.45, 2.6, 0.5, [[(c['capital'], {'size': 26}),
                                    ('  / ' + c['capital_period'], {'size': 11, 'font': T.LABEL_FONT})]],
         font=T.TITLE_FONT, color=T.WHITE, anchor=MID)
    for j, (pct, use) in enumerate(c['use']):
        y = 2.08 + j * 0.37
        text(s, 6.68, y, 0.55, 0.3, pct, size=11.5, font=T.TITLE_FONT, color=T.WHITE, anchor=MID)
        text(s, 7.25, y, 2.05, 0.3, use, size=10, color=T.LAV2, anchor=MID)
    rect(s, 0.6, 4.62, 8.8, 0.52, fill=T.PALEY, radius=0.14)
    text(s, 0.78, 4.62, 8.45, 0.52, c['cac'], size=9.5, font=T.LABEL_FONT, color=T.AMBER_DARK, anchor=MID,
         spacing=1.0)
    notes(s, c['notes'])


# ---------------------------------------------------------------- 12
def s_closing(s):
    c = K.CLOSING5
    title(s, c['title'], 0.6, 0.45, 8.0, 0.5, size=24)
    for i, (ic, head, detail) in enumerate(c['asks']):
        x = 0.6 + i * 4.5
        card(s, x, 1.08, 4.3, 0.7, radius=0.14)
        num_circle(s, i + 1, x + 0.15, 1.27, 0.32, size=10)
        text(s, x + 0.58, 1.13, 3.6, 0.26, head, size=11.5, font=T.TITLE_FONT, anchor=MID)
        text(s, x + 0.58, 1.39, 3.6, 0.36, detail, size=8.5, color=T.BODY, spacing=1.0)
    num_circle(s, 3, 0.6, 1.95, 0.32, size=10)
    text(s, 1.02, 1.93, 6.0, 0.36, c['pilot_title'], size=13, font=T.TITLE_FONT, color=T.DEEP, anchor=MID)
    sw, sg = 2.05, 0.2
    for i, (head, detail) in enumerate(c['steps']):
        x = 0.6 + i * (sw + sg)
        rect(s, x, 2.38, sw, 0.95, fill=T.LAV2 if i < 3 else T.PURPLE, radius=0.14)
        text(s, x + 0.14, 2.45, sw - 0.28, 0.2, head, size=8, font=T.LABEL_FONT, color=T.DEEP if i < 3 else T.LAV2)
        text(s, x + 0.14, 2.66, sw - 0.28, 0.64, detail, size=9, color=T.INK if i < 3 else T.WHITE, spacing=1.0)
        if i < 3:
            rect(s, x + sw + 0.05, 2.72, 0.1, 0.26, fill=T.PURPLE, shape=MSO_SHAPE.CHEVRON)
    card(s, 0.6, 3.45, 4.3, 0.95, radius=0.14)
    text(s, 0.76, 3.52, 4.0, 0.2, c['resources_title'], size=7.5, font=T.LABEL_FONT, color=T.PURPLE)
    for i, r in enumerate(c['resources']):
        x, y = 0.76 + (i % 2) * 2.05, 3.76 + (i // 2) * 0.3
        icon(s, 'check', 'purple', x, y + 0.05, 0.14)
        text(s, x + 0.2, y, 1.85, 0.26, r, size=9, anchor=MID)
    rect(s, 5.1, 3.45, 4.3, 0.95, fill=T.PALEY, radius=0.14)
    text(s, 5.26, 3.52, 4.0, 0.2, c['principle_title'], size=7.5, font=T.LABEL_FONT, color=T.AMBER_DARK)
    text(s, 5.26, 3.72, 4.0, 0.66, c['principle'], size=8.5, color=T.INK, spacing=1.02)
    rect(s, 0.6, 4.52, 7.5, 0.4, fill=T.PURPLE, radius=0.2)
    text(s, 0.6, 4.52, 7.5, 0.4, c['closing'], size=12, font=T.TITLE_FONT, color=T.WHITE, align=C, anchor=MID)
    people = '   ·   '.join(f'{n}, {r.split(" &")[0]}' for n, r in c['people'])
    box = text(s, 0.6, 4.98, 7.5, 0.24, [[(people + '   ·   ', {}), (c['website'], {'color': T.PURPLE})]], size=9,
               font=T.LABEL_FONT, align=C, anchor=MID)
    box.text_frame.paragraphs[0].runs[-1].hyperlink.address = 'https://www.scanvas.id/'
    _qr(s, 8.4, 4.42, 0.76)
    text(s, 8.25, 5.18, 1.06, 0.14, c['qr_label'], size=6.5, font=T.LABEL_FONT, color=T.PURPLE, align=C, anchor=MID)
    notes(s, c['notes'])


# ---------------------------------------------------------------- appendix
def s_financials(s):
    c = K.FINANCIALS5
    title(s, c['title'], 0.6, 0.45, 8.0, 0.5, size=24)
    x0, w = 0.6, 5.3
    card(s, x0, 1.1, w, 3.84, radius=0.14)
    for j, h in enumerate(c['head']):
        text(s, x0 + 0.18 if j == 0 else x0 + 2.95 + (j - 1) * 1.15, 1.17, 2.6 if j == 0 else 1.05, 0.2, h.upper(),
             size=7.5, font=T.LABEL_FONT, color=T.PURPLE, align=L if j == 0 else R)
    rh = 0.235
    for i, (lbl, a, b, style) in enumerate(c['rows']):
        y = 1.4 + i * rh
        if style in ('sum', 'cash'):
            rect(s, x0 + 0.08, y - 0.01, w - 0.16, rh + 0.02, fill=T.MINT if style == 'cash' else T.LAV, radius=0.06)
        else:
            rect(s, x0 + 0.18, y - 0.005, w - 0.36, 0.005, fill=T.BORDER)
        strong = style in ('sum', 'cash')
        text(s, x0 + (0.34 if style == 'sub' else 0.18), y, 2.7, rh, lbl, size=8.5,
             font=T.LABEL_FONT if strong else T.BODY_FONT, color=T.BODY if style == 'sub' else T.INK, anchor=MID)
        for j, v in enumerate((a, b)):
            text(s, x0 + 2.95 + j * 1.15, y, 1.05, rh, v, size=9 if strong else 8.5,
                 font=T.TITLE_FONT if strong else T.BODY_FONT,
                 color=T.GREEN if style == 'cash' else (T.DEEP if strong else T.INK), align=R, anchor=MID)
    rect(s, 6.1, 1.1, 3.3, 1.5, fill=T.MINT, radius=0.14)
    text(s, 6.26, 1.18, 3.0, 0.2, c['cash_title'], size=7.5, font=T.LABEL_FONT, color=T.GREEN)
    for i, t in enumerate(c['cash']):
        y = 1.42 + i * 0.38
        icon(s, 'check', 'green', 6.26, y + 0.04, 0.14)
        text(s, 6.46, y, 2.85, 0.36, t, size=8.5, color=T.INK, spacing=1.0)
    card(s, 6.1, 2.72, 3.3, 2.22, radius=0.14)
    text(s, 6.26, 2.8, 3.0, 0.2, c['assump_title'], size=7.5, font=T.LABEL_FONT, color=T.PURPLE)
    for i, a in enumerate(c['assumptions']):
        y = 3.04 + i * 0.37
        oval(s, 6.28, y + 0.09, 0.07, fill=T.PURPLE)
        text(s, 6.44, y, 2.88, 0.36, a, size=8.5, color=T.INK, spacing=1.0)
    text(s, 0.6, 5.0, 8.8, 0.28, c['basis'], size=7, color=T.BODY, spacing=1.0)
    notes(s, c['notes'])


def s_chaoyang(s):
    c = K.CHAOYANG5
    title(s, c['title'], 0.6, 0.45, 8.0, 0.5, size=24)
    rect(s, 0.6, 1.1, 8.8, 0.72, fill=T.LAV2, radius=0.16)
    text(s, 0.8, 1.16, 2.0, 0.2, c['objective_title'], size=7.5, font=T.LABEL_FONT, color=T.PURPLE)
    text(s, 0.8, 1.36, 8.4, 0.42, c['objective'], size=12.5, font=T.TITLE_FONT, color=T.INK, anchor=MID)
    for j, (head, items) in enumerate([(c['scope_title'], c['scope']), (c['measure_title'], c['measure'])]):
        x = 0.6 + j * 4.5
        card(s, x, 1.95, 4.3, 2.3, radius=0.16)
        text(s, x + 0.2, 2.05, 3.9, 0.2, head, size=7.5, font=T.LABEL_FONT, color=T.PURPLE)
        for i, it in enumerate(items):
            y = 2.32 + i * 0.31
            icon(s, 'check' if j == 0 else 'arrow-right', 'purple', x + 0.2, y + 0.06, 0.15)
            text(s, x + 0.44, y, 3.7, 0.28, it, size=10.5, anchor=MID)
    rect(s, 0.6, 4.38, 8.8, 0.76, fill=T.PURPLE, radius=0.16)
    text(s, 0.8, 4.45, 3.0, 0.2, c['gate_title'], size=7.5, font=T.LABEL_FONT, color=T.LAV2)
    text(s, 0.8, 4.64, 8.4, 0.44, c['gate'], size=13, font=T.TITLE_FONT, color=T.WHITE, anchor=MID)
    notes(s, c['notes'])
