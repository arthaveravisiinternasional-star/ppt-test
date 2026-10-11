"""Build the ITEC SEA 2026 7-minute pitch edition (v5): 12 main slides + appendix on the committee template.
Copy = founders' 12-slide spec (content_itec5.py); appendix reuses v4 slides. v1-v4 files and outputs are untouched.

    python build_deck_itec5.py                  # -> output/itec5/SCANVAS_ITEC_SEA2026_v5.pptx
    python build_deck_itec5.py --render 6 7     # previews in output/itec5/preview
    python build_deck_itec5.py --render --pdf   # final: embed fonts + PDF
"""
import argparse
import copy
import subprocess
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

import deck
from deck import content_itec5

sys.modules['deck.content'] = content_itec5   # swap copy before slide modules import it
deck.content = content_itec5

from deck import theme as T  # noqa: E402

# Typography of the committee template: Roboto Condensed (static cuts in assets/fonts, tools/make_fonts.py).
# Set before helpers is imported, because helper defaults read the body font at import time.
T.TITLE_FONT, T.LABEL_FONT, T.BODY_FONT = 'Roboto Condensed Bold', 'Roboto Condensed SemiBold', 'Roboto Condensed'
T.BOLD_FONTS = {T.TITLE_FONT}
T.FONT_ALIAS = {'Roboto Condensed Bold': 'Roboto Condensed'}   # written as family + bold, so it embeds

from deck import helpers  # noqa: E402
from deck import slides_dsc as D, slides_itec3 as I3, slides_itec4 as I4, slides_itec5 as I5  # noqa: E402

helpers.SCALE = 2.0                       # template canvas is 20 x 11.25 in; slide code uses the 10 x 5.625 grid
TEMPLATE = T.ASSETS / 'template' / 'itec_official.pptx'   # committee file (ITEC PITCH DECK.pptx), never edited
OUT_DIR = T.OUTPUT / 'itec5'
OUT = OUT_DIR / 'SCANVAS_ITEC_SEA2026_v5.pptx'

# (slide function, committee section label). The first entry reuses the template's own cover slide.
DECK = [(I5.s_cover, ''),                                       # 1 Cover            0:15
        (I5.s_problem, 'PROBLEM'),                              # 2 Problem          0:40
        (I5.s_solution, 'SOLUTION'),                            # 3 Solution         0:25
        (I5.s_demo, 'PRODUCT OR SERVICE · LIVE DEMO'),          # 4 Demo (video)     0:55
        (I5.s_traction, 'TRACTION'),                            # 5 Pilot            0:35
        (I5.s_market, 'TARGET MARKET · MARKET SIZE'),           # 6 TAM/SAM/beachhead 0:35
        (I5.s_competition, 'COMPETITORS · UNIQUE VALUE'),       # 7                  0:25
        (I5.s_business, 'BUSINESS OR REVENUE MODEL'),           # 8                  0:35
        (I5.s_team, 'MEET OUR TEAM'),                           # 9                  0:30
        (I5.s_roadmap, 'ROADMAP'),                              # 10                 0:30
        (I5.s_funds, 'THE ASK · ECONOMICS'),                    # 11 Economics + ask 0:40
        (I5.s_closing, 'ITEC FIT · CHAOYANG PILOT'),            # 12 ITEC/Chaoyang   0:35
        (I4.s_divider, ''),                                     # appendix for Q&A
        (I4.s_competition, 'APPENDIX · COMPETITOR MATRIX'),
        (I4.s_market_size, 'APPENDIX · MARKET EXPANSION'),
        (D.s_sources, 'APPENDIX · MARKET DATA'),
        (I3.s_pipeline, 'APPENDIX · TECHNOLOGY & CHILD DATA'),
        (I5.s_financials, 'APPENDIX · FINANCIAL PROJECTION'),
        (I5.s_chaoyang, 'APPENDIX · CHAOYANG PILOT')]


def _background(s):
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = helpers.rgb(T.CANVAS)


def build(out=OUT):
    prs = Presentation(str(TEMPLATE))
    cover, frame_src = prs.slides[0], prs.slides[1]
    frame = copy.deepcopy(next(sh for sh in frame_src.shapes if sh.name == 'Group 2')._element)   # blue/green bars
    for sh in list(cover.shapes):              # keep the cover's bars + ITEC logo, drop placeholder text
        if sh.shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
            sh._element.getparent().remove(sh._element)
    ids = prs.slides._sldIdLst
    for sid in list(ids)[1:]:                  # drop template slides 2-20 (their parts are not saved)
        prs.part.drop_rel(sid.rId)
        ids.remove(sid)
    blank = next(L for L in prs.slide_layouts if L.name == 'Blank')
    for i, (fn, label) in enumerate(DECK):
        if i == 0:
            s = cover
        else:
            s = prs.slides.add_slide(blank)
            s.shapes._spTree.insert(2, copy.deepcopy(frame))
        _background(s)
        fn(s)
        if label:
            I4.section(s, label)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    print('saved', out)
    return out


def powerpoint(out, mode, slides=()):
    cmd = ['powershell', '-NoProfile', '-ExecutionPolicy', 'Bypass', '-File', str(T.ROOT / 'tools' / 'render.ps1'),
           '-File', str(out), '-Mode', mode, '-Dir', str(OUT_DIR / 'preview')]
    if slides:
        cmd += ['-Slides', ','.join(map(str, slides))]
    subprocess.run(cmd, check=True)


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--render', nargs='*', type=int)
    ap.add_argument('--pdf', action='store_true')
    ap.add_argument('--out', help='output .pptx path relative to project root')
    a = ap.parse_args()
    out = build(T.ROOT / a.out if a.out else OUT)
    if a.render is not None:
        powerpoint(out, 'png', a.render)
    if a.pdf:
        powerpoint(out, 'embed')
        powerpoint(out, 'pdf')
