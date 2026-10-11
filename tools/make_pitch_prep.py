"""Printable DOCX + PDF of the ITEC pitch prep page (script, Q&A, numbers, answer rules).

Single source: docs/pitch_prep_itec5.html (the published artifact page). Edit that file, then:

    python tools/make_pitch_prep.py   # -> output/itec5/SCANVAS_ITEC_pitch_prep.docx + .pdf (A4, fonts embedded)

The PDF and font embedding go through Microsoft Word (COM), like the decks go through PowerPoint.
"""
import re
import subprocess
from pathlib import Path

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / 'docs' / 'pitch_prep_itec5.html'
OUT = ROOT / 'output' / 'itec5' / 'SCANVAS_ITEC_pitch_prep.docx'

DISPLAY, BODY = 'Fredoka', 'Varela Round'
INK, MUTED, PURPLE, DEEP = '172033', '526078', '7C3AED', '5B21B6'
S_INK, H_INK, WARN = '5B21B6', '1D4ED8', '92400E'
LAV, WARN_BG, OK_BG, OK_INK = 'F5F3FF', 'FEF3C7', 'DCFCE7', '15803D'


def rgb(h):
    return RGBColor.from_string(h)


P_AFTER_SHD = ('w:tabs', 'w:suppressAutoHyphens', 'w:kinsoku', 'w:wordWrap', 'w:overflowPunct', 'w:topLinePunct',
               'w:autoSpaceDE', 'w:autoSpaceDN', 'w:bidi', 'w:adjustRightInd', 'w:snapToGrid', 'w:spacing', 'w:ind',
               'w:contextualSpacing', 'w:mirrorIndents', 'w:suppressOverlap', 'w:jc', 'w:textDirection',
               'w:textAlignment', 'w:textboxTightWrap', 'w:outlineLvl', 'w:divId', 'w:cnfStyle', 'w:rPr',
               'w:sectPr', 'w:pPrChange')
TC_AFTER_SHD = ('w:noWrap', 'w:tcMar', 'w:textDirection', 'w:tcFitText', 'w:vAlign', 'w:hideMark')


def shade(el, fill):
    """Background fill for a paragraph (w:p) or table cell (w:tcPr), in schema order."""
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
    if hasattr(el, 'get_or_add_pPr'):
        el.get_or_add_pPr().insert_element_before(shd, *P_AFTER_SHD)
    else:
        el.insert_element_before(shd, *TC_AFTER_SHD)


def para(doc, size=11, color=INK, before=0, after=6, keep=False, line=1.15):
    p = doc.add_paragraph()
    f = p.paragraph_format
    f.space_before, f.space_after, f.line_spacing = Pt(before), Pt(after), line
    f.keep_with_next = keep
    p._defaults = {'size': size, 'color': color}
    return p


def run(p, t, bold=False, color=None, font=None, size=None, highlight=False, italic=False):
    r = p.add_run(t)
    d = getattr(p, '_defaults', {'size': 11, 'color': INK})
    r.font.name = font or BODY
    r._element.rPr.rFonts.set(qn('w:eastAsia'), font or BODY)
    r.font.size = Pt(size or d['size'])
    r.font.color.rgb = rgb(color or d['color'])
    r.font.bold, r.font.italic = bold, italic
    if highlight:
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return r


def inline(p, node, **kw):
    """Walk an HTML element and add runs, keeping bold, fill-in highlights, speakers and stage directions."""
    for ch in node.children:
        if isinstance(ch, NavigableString):
            t = re.sub(r'\s+', ' ', str(ch))
            if t:
                run(p, t, **kw)
            continue
        cls = ch.get('class', [])
        if 'fill' in cls:
            run(p, ch.get_text(' ', strip=True), highlight=True, color=WARN, size=kw.get('size'))
        elif 'sp' in cls:
            run(p, ch.get_text(strip=True) + '  ', bold=True, font=DISPLAY,
                color=S_INK if 's' in cls else H_INK, size=kw.get('size'))
        elif 'stage' in cls:
            run(p, ch.get_text(' ', strip=True), bold=True, color=DEEP, font=DISPLAY, size=kw.get('size'))
        elif 'by' in cls:
            run(p, '   ' + ch.get_text(strip=True), color=MUTED, size=9)
        elif ch.name in ('b', 'strong'):
            inline(p, ch, **{**kw, 'bold': True})
        else:
            inline(p, ch, **kw)


def heading(doc, text, size, color=INK, before=0, after=6, eyebrow=None, page_break=False):
    first = None
    if eyebrow:
        e = first = para(doc, size=8.5, color=PURPLE, before=before, after=2, keep=True)
        run(e, eyebrow.upper(), bold=True, font=DISPLAY)
        before = 0
    p = para(doc, size=size, color=color, before=before, after=after, keep=True, line=1.0)
    (first or p).paragraph_format.page_break_before = page_break
    run(p, text, bold=True, font=DISPLAY)
    return p


def page_number_footer(section):
    p = section.footer.paragraphs[0]
    p._defaults = {'size': 8, 'color': MUTED}
    run(p, 'Scanvas · ITEC SEA 2026 pitch prep · page ')
    r = p.add_run()
    for kind, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
        if kind:
            el = OxmlElement('w:fldChar'); el.set(qn('w:fldCharType'), kind)
        else:
            el = OxmlElement('w:instrText'); el.set(qn('xml:space'), 'preserve'); el.text = txt
        r._element.append(el)
    r.font.size, r.font.color.rgb, r.font.name = Pt(8), rgb(MUTED), BODY


def table(doc, head, rows, widths, value_col=0, status_col=None):
    t = doc.add_table(rows=1, cols=len(head))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    for i, h in enumerate(head):
        c = t.rows[0].cells[i]
        shade(c._tc.get_or_add_tcPr(), LAV)
        p = c.paragraphs[0]; p._defaults = {'size': 8.5, 'color': PURPLE}
        run(p, h.upper(), bold=True, font=DISPLAY)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            p = cells[i].paragraphs[0]; p._defaults = {'size': 10, 'color': INK}
            if i == status_col:
                shade(cells[i]._tc.get_or_add_tcPr(), OK_BG if v in ('Fakta', 'Data publik') else WARN_BG)
                run(p, v, bold=True, font=DISPLAY, size=8.5, color=OK_INK if v in ('Fakta', 'Data publik') else WARN)
            else:
                run(p, v, bold=i == value_col, font=DISPLAY if i == value_col else BODY)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    # light borders
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        b = OxmlElement(f'w:{side}')
        b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '4'); b.set(qn('w:color'), 'D9DEEA')
        borders.append(b)
    t._tbl.tblPr.insert_element_before(borders, 'w:shd', 'w:tblLayout', 'w:tblCellMar', 'w:tblLook',
                                        'w:tblCaption', 'w:tblDescription')
    para(doc, after=4)
    return t


def build():
    soup = BeautifulSoup(SRC.read_text(encoding='utf-8'), 'html.parser')
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)
    page_number_footer(sec)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = BODY, Pt(11)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), BODY)

    # title block
    top = soup.select_one('header.top')
    heading(doc, top.h1.get_text(strip=True), 26, eyebrow=top.select_one('.eyebrow').get_text(strip=True), after=6)
    p = para(doc, size=10.5, color=MUTED, after=6)
    inline(p, top.select_one('.lede'))
    p = para(doc, size=10, after=14)
    for i, pill in enumerate(top.select('.pill')):
        b = pill.b.get_text(strip=True)
        rest = pill.get_text(' ', strip=True)[len(b):].strip()
        run(p, ('     ' if i else '') + b + ' ', bold=True, font=DISPLAY, color=PURPLE)
        run(p, rest)

    # 1 script
    s = soup.select_one('#script')
    heading(doc, s.h2.get_text(strip=True), 20, eyebrow=s.select_one('.eyebrow').get_text(strip=True), after=4)
    p = para(doc, size=10.5, color=MUTED, after=4); inline(p, s.select_one('.lede'))
    p = para(doc, size=9.5, after=4)
    run(p, 'S', bold=True, font=DISPLAY, color=S_INK); run(p, ' = Salsabila     ')
    run(p, 'H', bold=True, font=DISPLAY, color=H_INK); run(p, ' = Haikal     ')
    run(p, 'Highlight kuning = isi dulu (fakta yang hanya kalian tahu).', color=WARN)
    for art in s.select('article.slide'):
        head = art.select_one('.slide-head')
        num, name = head.select_one('.num').get_text(strip=True), head.h3.get_text(strip=True)
        whos = [w.get_text(strip=True) for w in head.select('.who')]
        time = head.select_one('.time').get_text(' ', strip=True)
        p = para(doc, size=14, before=12, after=3, keep=True, line=1.0)
        run(p, f'{num}  ', bold=True, font=DISPLAY, color=PURPLE)
        run(p, name, bold=True, font=DISPLAY)
        run(p, '     ' + ' + '.join(whos), bold=True, font=DISPLAY, size=10,
            color=H_INK if whos == ['H'] else S_INK)
        run(p, '   ·   ' + time, size=9.5, color=MUTED)
        for sp in art.select('.say p'):
            q = para(doc, size=12.5, after=5, line=1.3, keep=True)   # keep one slide on one page
            inline(q, sp)
        cue = art.select_one('.cue')
        if cue:
            q = para(doc, size=9.5, color=MUTED, before=2, after=4)
            shade(q._p, LAV)
            run(q, 'Catatan: ', bold=True, font=DISPLAY, color=PURPLE)
            inline(q, cue)

    # 2 Q&A
    s = soup.select_one('#qa')
    heading(doc, s.h2.get_text(strip=True), 20, eyebrow=s.select_one('.eyebrow').get_text(strip=True), after=4,
            page_break=True)
    p = para(doc, size=10.5, color=MUTED, after=6); inline(p, s.select_one('.lede'))
    n = 0
    for cat in s.select('.cat'):
        h3 = cat.h3
        small = h3.small.get_text(strip=True) if h3.small else ''
        title = h3.get_text(' ', strip=True).replace(small, '').strip()
        p = para(doc, size=14, color=DEEP, before=12, after=4, keep=True, line=1.0)
        run(p, title, bold=True, font=DISPLAY)
        if small:
            run(p, '   ' + small, size=9.5, color=MUTED)
        for d in cat.select('details.q'):
            n += 1
            p = para(doc, size=11.5, before=6, after=2, keep=True, line=1.1)
            run(p, f'Q{n}.  ', bold=True, font=DISPLAY, color=PURPLE)
            qt = d.select_one('.qtext')
            who = qt.select_one('.by').extract().get_text(strip=True)
            run(p, qt.get_text(' ', strip=True), bold=True, font=DISPLAY)
            run(p, f'   [{who}]', size=9, color=MUTED)
            for el in d.select('.ans > p'):
                cls = el.get('class', [])
                if 'a' in cls:
                    q = para(doc, size=11, after=3, line=1.2, keep=el.find_next_sibling('p') is not None)
                    q.paragraph_format.left_indent = Cm(0.6)
                    inline(q, el)
                elif 'tip' in cls:
                    q = para(doc, size=9.5, color=MUTED, after=3)
                    q.paragraph_format.left_indent = Cm(0.6)
                    inline(q, el)
                elif 'avoid' in cls:
                    q = para(doc, size=9.5, color=WARN, after=3)
                    q.paragraph_format.left_indent = Cm(0.6)
                    shade(q._p, WARN_BG)
                    inline(q, el)

    # 3 numbers
    s = soup.select_one('#numbers')
    heading(doc, s.h2.get_text(strip=True), 20, eyebrow=s.select_one('.eyebrow').get_text(strip=True), after=4,
            page_break=True)
    p = para(doc, size=10.5, color=MUTED, after=8); inline(p, s.select_one('.lede'))
    tables = s.select('table')
    rows = [[td.get_text(' ', strip=True) for td in tr.select('td')] for tr in tables[0].select('tbody tr')]
    table(doc, [th.get_text(strip=True) for th in tables[0].select('th')], rows, [3.2, 11.0, 2.8], status_col=2)
    rows = [[td.get_text(' ', strip=True) for td in tr.select('td')] for tr in tables[1].select('tbody tr')]
    table(doc, [th.get_text(strip=True) for th in tables[1].select('th')], rows, [4.6, 12.4])

    # 4 rules
    s = soup.select_one('#rules')
    heading(doc, s.h2.get_text(strip=True), 20, eyebrow=s.select_one('.eyebrow').get_text(strip=True), before=10,
            after=6)
    for li in s.select('li'):
        p = para(doc, size=11, after=4, line=1.2)
        p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Cm(0.6), Cm(-0.4)
        run(p, '•  ', color=PURPLE, bold=True)
        inline(p, li)

    for p in doc.paragraphs:                    # collapse edge spaces left by HTML whitespace
        if p.runs:
            p.runs[0].text = p.runs[0].text.lstrip()
            p.runs[-1].text = p.runs[-1].text.rstrip()
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print('saved', OUT, '|', n, 'questions')
    return OUT


def word_finish(path):
    """Embed fonts into the DOCX and export the PDF through Word."""
    ps = f'''
$w = New-Object -ComObject Word.Application; $w.Visible = $false
$d = $w.Documents.Open('{path}')
$d.EmbedTrueTypeFonts = $true; $d.SaveSubsetFonts = $true
$d.Save()
$d.ExportAsFixedFormat('{path.with_suffix(".pdf")}', 17)
$d.Close(); $w.Quit()
'''
    subprocess.run(['powershell', '-NoProfile', '-Command', ps], check=True)
    print('saved', path.with_suffix('.pdf'))


if __name__ == '__main__':
    word_finish(build())
