"""Build the SCANVAS investor deck (October 2026): 13 main slides + 7 appendix, Freesia template.
Branch of the DSC S17 deck, updated with ITEC v5, the new course roadmap, and the founders' CVs.

    python build_deck_investor.py                  # -> output/investor/SCANVAS_Investor_Oct2026.pptx
    python build_deck_investor.py --render 6 12    # previews in output/investor/preview
    python build_deck_investor.py --render --pdf   # final: embed fonts + PDF
"""
import argparse
import subprocess
import sys

from pptx import Presentation

import deck
from deck import content_investor

sys.modules['deck.content'] = content_investor   # swap copy before slide modules import it
deck.content = content_investor

from deck import recolor, slides as S, slides_dsc as D, slides_itec3 as I3, slides_investor as V  # noqa: E402
from deck import theme as T  # noqa: E402

OUT_DIR = T.OUTPUT / 'investor'
OUT = OUT_DIR / 'SCANVAS_Investor_Oct2026.pptx'

# (slide function, Freesia layout) in DSC guideline order
DECK = [(S.s1_cover, 'TITLE'),                      # 1 Scanvas
        (V.s_problem, 'TITLE_ONLY_3_2'),            # 2 Customer problem
        (V.s_product, 'TITLE_ONLY'),                # 3 Product experience + QR
        (V.s_traction, 'TITLE_ONLY_3_1_2'),         # 4 Traction & pilot learning
        (V.s_market, 'TITLE_AND_BODY_1_1'),         # 5 Customers & market
        (V.s_business, 'TITLE_AND_BODY'),           # 6 Business model & unit economics
        (V.s_competition, 'TITLE_ONLY'),            # 7 Competition & differentiation
        (V.s_gtm, 'TITLE_ONLY'),                    # 8 Go-to-market & repeat purchase
        (V.s_team, 'TITLE_ONLY_2_1'),               # 9 Founding team
        (V.s_plan, 'BLANK_1_1'),                    # 10 12-month execution plan
        (V.s_vision, 'TITLE_ONLY_3_2'),             # 11 Five-year growth vision
        (V.s_fin, 'TITLE_AND_BODY'),                # 12 Financial projections
        (V.s_ask, 'TITLE_ONLY_3_1_2'),              # 13 Investment ask + contact
        (V.s_appx_fin, 'TITLE_ONLY'),               # A Unit costs, volumes, annual cash (capex + working capital)
        (V.s_appx_cash, 'TITLE_ONLY'),              # B Monthly cash, months 1-24 (sizes the year-2 raise)
        (V.s_appx_sales, 'TITLE_ONLY'),             # C Sales engine: target -> channel -> conversion -> cost -> capacity
        (D.s_sources, 'TITLE_ONLY_3_1_2'),          # D Market size sources
        (I3.s_competition, 'TITLE_ONLY'),           # E Competitor evidence
        (I3.s_pipeline, 'TITLE_ONLY_3_2'),          # F Technology & child data
        (V.s_appx_course, 'TITLE_ONLY')]            # G Course MVP as a learning product


def build(out=OUT):
    prs = Presentation(str(T.TEMPLATE))
    recolor.apply(prs)
    layouts = {L.name: L for m in prs.slide_masters for L in m.slide_layouts}
    ids = prs.slides._sldIdLst
    for sid in list(ids):
        prs.part.drop_rel(sid.rId)
        ids.remove(sid)
    for fn, layout in DECK:
        s = prs.slides.add_slide(layouts[layout])
        for ph in list(s.placeholders):
            if ph.placeholder_format.idx != 0:
                ph._element.getparent().remove(ph._element)
        fn(s)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    print('saved', out)
    return out


def powerpoint(out, mode, slides=()):
    cmd = ['powershell', '-NoProfile', '-ExecutionPolicy', 'Bypass', '-File', str(T.ROOT / 'tools' / 'render.ps1'),
           '-File', str(out), '-Mode', mode, '-Dir', str(OUT_DIR / 'preview')]
    if slides:
        cmd += ['-Slides', ','.join(map(str, slides))]
    subprocess.run(cmd, check=True)


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--render', nargs='*', type=int)
    ap.add_argument('--pdf', action='store_true')
    ap.add_argument('--out', help='output .pptx path relative to project root')
    a = ap.parse_args()
    out = build(T.ROOT / a.out if a.out else OUT)
    if a.render is not None:
        powerpoint(out, 'png', a.render)
    if a.pdf:
        powerpoint(out, 'embed')
        powerpoint(out, 'pdf')
