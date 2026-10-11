"""Build the ITEC SEA 2026 official-template edition (v4) on the committee's pitch-deck template.
Copy = ITEC v3 final, mapped onto the committee's sections. v1/v2/v3 files and outputs are untouched.

    python build_deck_itec4.py                  # -> output/itec4/SCANVAS_ITEC_SEA2026_v4.pptx
    python build_deck_itec4.py --render 6 7     # previews in output/itec4/preview
    python build_deck_itec4.py --render --pdf   # final: embed fonts + PDF
"""
import argparse
import copy
import subprocess
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

import deck
from deck import content_itec4

sys.modules['deck.content'] = content_itec4   # swap copy before slide modules import it
deck.content = content_itec4

from deck import theme as T  # noqa: E402

# Typography of the committee template: Roboto Condensed (static cuts in assets/fonts, tools/make_fonts.py).
# Set before helpers is imported, because helper defaults read the body font at import time.
T.TITLE_FONT, T.LABEL_FONT, T.BODY_FONT = 'Roboto Condensed Bold', 'Roboto Condensed SemiBold', 'Roboto Condensed'
T.BOLD_FONTS = {T.TITLE_FONT}
T.FONT_ALIAS = {'Roboto Condensed Bold': 'Roboto Condensed'}   # written as family + bold, so it embeds

from deck import helpers  # noqa: E402
from deck import slides_dsc as D, slides_itec3 as I3, slides_itec4 as I4  # noqa: E402

helpers.SCALE = 2.0                       # template canvas is 20 x 11.25 in; slide code uses the 10 x 5.625 grid
TEMPLATE = T.ASSETS / 'template' / 'itec_official.pptx'   # committee file (ITEC PITCH DECK.pptx), never edited
OUT_DIR = T.OUTPUT / 'itec4'
OUT = OUT_DIR / 'SCANVAS_ITEC_SEA2026_v4.pptx'

# (slide function, committee section label). The first entry reuses the template's own cover slide.
DECK = [(I4.s_cover, ''),                                       # 1 Cover (template slide 1 + ITEC logo)
        (I3.s_problem, 'PROBLEM'),                              # 2
        (I3.s_solution, 'SOLUTION'),                            # 3
        (I4.s_product, 'PRODUCT OR SERVICE'),                   # 4
        (I4.s_demo, 'PRODUCT OR SERVICE · LIVE DEMO'),          # 5 video of the digital trial
        (I4.s_about, 'ABOUT OUR COMPANY'),                      # 6
        (I3.s_traction, 'ABOUT OUR COMPANY · TRACTION'),        # 7
        (I4.s_target, 'TARGET MARKET'),                         # 8
        (I4.s_market_size, 'MARKET SIZE / POTENTIAL'),          # 9
        (I4.s_competition, 'COMPETITORS'),                      # 10
        (I4.s_uvp, 'UNIQUE VALUE PROPOSITION'),                 # 11
        (I3.s_tech, 'COMPETITIVE ADVANTAGE'),                   # 12
        (I3.s_culture, 'DIGITAL CULTURE TRACK FIT'),            # 13
        (I3.s_business, 'BUSINESS OR REVENUE MODEL'),           # 14
        (I3.s_repeat, 'BUSINESS MODEL · REPEAT VALUE'),         # 15
        (I4.s_team, 'MEET OUR TEAM'),                           # 16
        (I3.s_next, 'ROADMAP'),                                 # 17
        (I4.s_ask, 'THE ASK'),                                  # 18
        (I4.s_funds, 'THE ASK · USE OF FUNDS'),                 # 19 founders' 12-month model (Rp200M)
        (I4.s_exit, 'EXIT STRATEGY'),                           # 20 indicative
        (I4.s_contact, ''),                                     # 21 Contact
        (I4.s_divider, ''),                                     # 22 Supporting slides
        (D.s_sources, 'OTHER INFO · MARKET DATA'),              # 23
        (I3.s_pipeline, 'OTHER INFO · TECHNOLOGY & CHILD DATA')]  # 24


def _background(s):
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = helpers.rgb(T.CANVAS)


def build(out=OUT):
    prs = Presentation(str(TEMPLATE))
    cover, frame_src = prs.slides[0], prs.slides[1]
    frame = copy.deepcopy(next(sh for sh in frame_src.shapes if sh.name == 'Group 2')._element)   # blue/green bars
    for sh in list(cover.shapes):              # keep the cover's bars + ITEC logo, drop placeholder text
        if sh.shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
            sh._element.getparent().remove(sh._element)
    ids = prs.slides._sldIdLst
    for sid in list(ids)[1:]:                  # drop template slides 2-20 (their parts are not saved)
        prs.part.drop_rel(sid.rId)
        ids.remove(sid)
    blank = next(L for L in prs.slide_layouts if L.name == 'Blank')
    for i, (fn, label) in enumerate(DECK):
        if i == 0:
            s = cover
        else:
            s = prs.slides.add_slide(blank)
            s.shapes._spTree.insert(2, copy.deepcopy(frame))
        _background(s)
        fn(s)
        if label:
            I4.section(s, label)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    print('saved', out)
    return out


def powerpoint(out, mode, slides=()):
    cmd = ['powershell', '-NoProfile', '-ExecutionPolicy', 'Bypass', '-File', str(T.ROOT / 'tools' / 'render.ps1'),
           '-File', str(out), '-Mode', mode, '-Dir', str(OUT_DIR / 'preview')]
    if slides:
        cmd += ['-Slides', ','.join(map(str, slides))]
    subprocess.run(cmd, check=True)


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--render', nargs='*', type=int)
    ap.add_argument('--pdf', action='store_true')
    ap.add_argument('--out', help='output .pptx path relative to project root')
    a = ap.parse_args()
    out = build(T.ROOT / a.out if a.out else OUT)
    if a.render is not None:
        powerpoint(out, 'png', a.render)
    if a.pdf:
        powerpoint(out, 'embed')
        powerpoint(out, 'pdf')
