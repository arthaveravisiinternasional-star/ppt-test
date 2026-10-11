"""Build the Antler (pre-seed VC) edition: same template, look, helpers and assets; investor slide order,
copy from deck/content_antler.py, plus market and ask slides from deck/slides_antler.py.
GESAwards (build_deck.py) and ITEC (build_deck_itec.py) editions are not touched.

    python build_deck_antler.py                 # -> output/antler/SCANVAS_Antler_PreSeed.pptx
    python build_deck_antler.py --render 7 12   # previews in output/antler/preview
    python build_deck_antler.py --render --pdf  # final: embed fonts + PDF
"""
import argparse
import subprocess
import sys

from pptx import Presentation

import deck
from deck import content_antler

sys.modules['deck.content'] = content_antler   # swap copy before slide modules import it
deck.content = content_antler

from deck import recolor, slides as S, slides_antler as A  # noqa: E402
from deck import theme as T  # noqa: E402

OUT_DIR = T.OUTPUT / 'antler'
OUT = OUT_DIR / 'SCANVAS_Antler_PreSeed.pptx'

# (slide function, Freesia layout) in investor-story order
DECK = [(S.s1_cover, 'TITLE'),
        (S.s2_problem, 'TITLE_ONLY_3_2'),
        (S.s4_solution, 'TITLE_ONLY'),
        (S.s5_product, 'TITLE_ONLY_3_1_1_1'),
        (S.s6_magic, 'CAPTION_ONLY'),
        (S.s7_different, 'TITLE_ONLY_2_1'),
        (A.s_market, 'TITLE_AND_BODY_1_1'),
        (S.s10_market, 'TITLE_AND_BODY'),
        (S.s11_validation, 'TITLE_ONLY_3_1_2'),
        (S.s8_roadmap, 'BLANK_1_1'),
        (S.s12_team, 'SECTION_HEADER_1_1'),
        (A.s_ask, 'TITLE_ONLY_3_2')]


def build(out=OUT):
    prs = Presentation(str(T.TEMPLATE))
    recolor.apply(prs)
    layouts = {L.name: L for m in prs.slide_masters for L in m.slide_layouts}
    ids = prs.slides._sldIdLst
    for sid in list(ids):
        prs.part.drop_rel(sid.rId)
        ids.remove(sid)
    for fn, layout in DECK:
        s = prs.slides.add_slide(layouts[layout])
        for ph in list(s.placeholders):
            if ph.placeholder_format.idx != 0:
                ph._element.getparent().remove(ph._element)
        fn(s)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    print('saved', out)
    return out


def powerpoint(out, mode, slides=()):
    cmd = ['powershell', '-NoProfile', '-ExecutionPolicy', 'Bypass', '-File', str(T.ROOT / 'tools' / 'render.ps1'),
           '-File', str(out), '-Mode', mode, '-Dir', str(OUT_DIR / 'preview')]
    if slides:
        cmd += ['-Slides', ','.join(map(str, slides))]
    subprocess.run(cmd, check=True)


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--render', nargs='*', type=int)
    ap.add_argument('--pdf', action='store_true')
    ap.add_argument('--out', help='output .pptx path relative to project root')
    a = ap.parse_args()
    out = build(T.ROOT / a.out if a.out else OUT)
    if a.render is not None:
        powerpoint(out, 'png', a.render)
    if a.pdf:
        powerpoint(out, 'embed')
        powerpoint(out, 'pdf')
